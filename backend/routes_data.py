import io
import re
import time
import uuid
from datetime import date, datetime, timedelta
from typing import Optional, Literal

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import String, case, cast, func, insert, or_, and_, desc, asc, text
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session, joinedload

import customer_tables
import models, schemas
from activity import log_activity
from auth import get_current_user, require_admin, require_staff, require_permission, get_user_permissions
from database import get_db, SessionLocal

# Every route here needs a signed-in user.
router = APIRouter(prefix="/api", dependencies=[Depends(get_current_user)])

Device, Customer = models.Device, models.Customer
ACTIVE = Device.deleted_at.is_(None)
BASE_COLUMNS = ["customer_name", "device_sr_no", "imei", "iccid", "msisdn1", "msisdn2", "act_date", "expiry_date",
                "dispatched_date", "model", "sim_provider", "activation_plan", "firmware", "flash", "sim_type"]
# Plain-text device fields that can be copied straight from an upload / edit form
TEXT_FIELDS = ["device_sr_no", "iccid", "msisdn1", "msisdn2", "model", "sim_provider", "activation_plan",
               "firmware", "flash"]

FIELD_MAX_LEN = {
    "imei": 20,
    "device_sr_no": 50,
    "iccid": 30,
    "msisdn1": 20,
    "msisdn2": 20,
    "model": 100,
    "sim_provider": 100,
    "activation_plan": 100,
    "firmware": 100,
    "flash": 100,
    "sim_type": 20,
    "customer_name": 150
}


def clean_str(val) -> Optional[str]:
    if val is None:
        return None
    s = str(val).strip()
    return s if s != "" else None


def clean_iccid(val: Optional[str]) -> str:
    if not val:
        return ""
    s = str(val).strip()
    if re.search(r"e\+", s, re.I):
        try:
            s = str(int(float(s)))
        except Exception:
            pass
    if s.endswith(".0"):
        s = s[:-2]
    s = re.sub(r"[\s\-_]+", "", s)
    return s.strip().upper()


def normalize_date_input(val):
    return schemas._to_date(val)


def escape_like(s: str) -> str:
    """So a search for 50% or a_b means those exact characters, not wildcards."""
    return s.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def classify_sim_type(iccid: Optional[str]) -> str:
    if iccid and str(iccid).strip():
        return "E-SIM"
    return "Plastic SIM"


def register_columns(db: Session, names):
    known = {n for (n,) in db.query(models.ColumnRegistry.name)}
    for n in set(names) - known:
        db.add(models.ColumnRegistry(name=n))


def clean_extra(extra):
    # keep it small and safe: short names, string values, no clash with built-in columns
    out = {}
    for k, v in (extra or {}).items():
        if v is None:
            continue
        k = str(k).strip()[:100]
        if k and k.lower().replace(' ', '_') not in BASE_COLUMNS and len(out) < 100:
            out[k] = str(v).strip()[:500]
    return out


def merge_extra(old, new):
    """Edit form: a filled value sets the column, a blank one clears it."""
    merged = dict(old or {})
    for k, v in (new or {}).items():
        k = str(k).strip()[:100]
        if not k or k.lower().replace(' ', '_') in BASE_COLUMNS:
            continue
        if v is None or str(v).strip() == "":
            merged.pop(k, None)
        else:
            merged[k] = str(v).strip()[:500]
    return merged


def compute_expiry_status(exp_date: Optional[date], window_days: int = 30) -> str:
    if not exp_date:
        return "Unknown"
    today = date.today()
    if exp_date < today:
        return "Expired"
    elif exp_date <= today + timedelta(days=window_days):
        return "Expiring Soon"
    else:
        return "Active"


def out(d: Device) -> dict:
    saved_stype = (d.sim_type or "").strip()
    stype = saved_stype if saved_stype else ("E-SIM" if d.iccid and str(d.iccid).strip() else "Plastic SIM")
    return {
        "id": d.id, "customer_id": d.customer_id,
        "customer_name": d.customer.name if d.customer else None,
        "device_sr_no": d.device_sr_no, "imei": d.imei, "iccid": d.iccid,
        "msisdn1": d.msisdn1, "msisdn2": d.msisdn2,
        "act_date": d.act_date, "expiry_date": d.expiry_date, "dispatched_date": d.dispatched_date,
        "model": d.model, "sim_provider": d.sim_provider, "activation_plan": d.activation_plan,
        "firmware": d.firmware, "flash": d.flash, "sim_type": stype,
        "extra": d.extra or {},
        "days_left": (d.expiry_date - date.today()).days if d.expiry_date else None,
        "expiry_status": compute_expiry_status(d.expiry_date, 30),
        "added_by": d.added_by, "deleted_at": d.deleted_at,
        "deleted_by": d.deleted_by, "delete_reason": d.delete_reason,
    }


def customer_by_name(db: Session, name: str) -> Customer:
    raw = (name or "").strip()
    clean_name = re.sub(r"[^a-zA-Z0-9_]+", "_", raw).strip("_") or raw
    
    # 1. Check if customer exists by exact name, clean name, or normalized name (ignoring spaces/underscores)
    c = db.query(Customer).filter(
        or_(
            func.lower(Customer.name) == raw.lower(),
            func.lower(Customer.name) == clean_name.lower(),
            func.replace(func.lower(Customer.name), '_', ' ') == func.replace(raw.lower(), '_', ' '),
            func.replace(func.lower(Customer.name), '_', ' ') == func.replace(clean_name.lower(), '_', ' ')
        )
    ).first()
    
    if c and c.deleted_at:  # adding devices to a deleted customer's name brings the customer back
        c.deleted_at = c.deleted_by = None
        
    if not c:
        try:
            c = Customer(name=clean_name[:150])
            db.add(c)
            db.flush()
        except Exception:
            db.rollback()
            c = db.query(Customer).filter(
                or_(
                    func.lower(Customer.name) == raw.lower(),
                    func.lower(Customer.name) == clean_name.lower(),
                    func.replace(func.lower(Customer.name), '_', ' ') == func.replace(raw.lower(), '_', ' '),
                    func.replace(func.lower(Customer.name), '_', ' ') == func.replace(clean_name.lower(), '_', ' ')
                )
            ).first()
            if not c:
                raise
        try:
            if getattr(getattr(db, "bind", None), "dialect", None) and getattr(db.bind.dialect, "name", "") == "sqlite":
                customer_tables.ensure_table(db.connection(), c.id, c.name)
            else:
                with customer_tables.engine.begin() as conn:
                    customer_tables.ensure_table(conn, c.id, c.name)
        except Exception as e:
            print(f"Customer table create failed for {c.name}: {e}")
    return c


def active_by_imei(db: Session, imei: str) -> Device:
    d = (db.query(Device).options(joinedload(Device.customer))
         .filter(Device.imei == imei.strip(), ACTIVE).first())
    if not d:
        raise HTTPException(404, "No device found with that IMEI.")
    return d


def trash(d: Device, user: models.User, reason: Optional[str]):
    d.deleted_at = datetime.utcnow()
    d.deleted_by = user.user_id
    d.delete_reason = reason or "Deleted by user"


def in_chunks(db: Session, column, values, size: int = 2000):
    """Yields (column value) rows for `column IN values`, a slice at a time, so 80,000-row files never build one giant IN()."""
    values = list(values)
    for i in range(0, len(values), size):
        yield from db.query(column).filter(column.in_(values[i:i + size]))


EXPORT_COLS = [("customer_name", "Customer Name"), ("device_sr_no", "Device SR No"), ("imei", "IMEI"), ("iccid", "ICCID"),
               ("msisdn1", "MSISDN1"), ("msisdn2", "MSISDN2"), ("act_date", "Act Date"), ("expiry_date", "Expiry Date"),
               ("dispatched_date", "Dispatched"), ("model", "Model"), ("sim_provider", "SIM Provider"),
               ("activation_plan", "Activation Plan"), ("firmware", "Firmware"), ("flash", "Flash"), ("sim_type", "SIM Type")]


def make_excel(title: str, headers: list[str], rows: list[list]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (re.sub(r'[\\/*?:\[\]]', '_', title or "Report"))[:31]

    header_fill = PatternFill(start_color="14213D", end_color="14213D", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    zebra = PatternFill(start_color="F9FAFD", end_color="F9FAFD", fill_type="solid")
    border_side = Side(style='thin', color='D5D9E6')
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    fancy = len(rows) <= 5000  # cell-by-cell styling is too slow for tens of thousands of rows
    widths = [len(str(h)) for h in headers]
    for r_idx, r in enumerate(rows, start=2):
        vals = [str(v) if v is not None else "" for v in r]
        ws.append(vals)
        for i, v in enumerate(vals[:len(widths)]):
            if len(v) > widths[i]:
                widths[i] = len(v)
        if fancy:
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if r_idx % 2 == 1:
                    cell.fill = zebra

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 3, 12), 60)
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------- dashboard ----------
@router.get("/stats")
def stats(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    count = lambda *f: db.query(func.count(Device.id)).filter(*f).scalar()
    return {
        "customers": db.query(func.count(Customer.id)).filter(Customer.deleted_at.is_(None)).scalar(),
        "devices": count(ACTIVE),
        "added_by_you": count(ACTIVE, Device.added_by == user.id),
        "columns": len(BASE_COLUMNS) + db.query(func.count(models.ColumnRegistry.name)).scalar(),
        "deleted": count(Device.deleted_at.isnot(None)),
        "expiring_30_days": count(ACTIVE, Device.expiry_date <= date.today() + timedelta(days=30)),
        "errored_records": db.query(func.count(models.ErroredRecord.id)).scalar() or 0,
    }


@router.get("/columns", dependencies=[Depends(require_permission("view"))])
def columns(db: Session = Depends(get_db)):
    """Extra column names found in uploads, shown after the built-in columns."""
    return [n for (n,) in db.query(models.ColumnRegistry.name).order_by(models.ColumnRegistry.name)]


# ---------- customers ----------
@router.get("/customers", dependencies=[Depends(require_permission("view"))])
def customers(db: Session = Depends(get_db)):
    soon = date.today() + timedelta(days=30)
    rows = (db.query(Customer.id, Customer.name, Customer.avatar, func.count(Device.id),
                     func.sum(case((Device.expiry_date <= soon, 1), else_=0)))
            .outerjoin(Device, (Device.customer_id == Customer.id) & ACTIVE)
            .filter(Customer.deleted_at.is_(None))
            .group_by(Customer.id, Customer.name, Customer.avatar).order_by(Customer.name).all())
    return [{"id": i, "name": n, "avatar": a, "devices": c, "expiring_or_expired": int(e or 0)} for i, n, a, c, e in rows]


@router.post("/customers", status_code=201, dependencies=[Depends(require_permission("add"))])
def add_customer(data: schemas.CustomerIn, db: Session = Depends(get_db),
                 user: models.User = Depends(get_current_user)):
    c = customer_by_name(db, data.name.strip())
    log_activity(db, "customer_add", f"Customer {c.name}", user=user)
    db.commit()
    try:  # give the new customer its own (empty) table right away
        if getattr(getattr(db, "bind", None), "dialect", None) and getattr(db.bind.dialect, "name", "") == "sqlite":
            customer_tables.ensure_table(db.connection(), c.id, c.name)
            db.commit()
        else:
            with customer_tables.engine.begin() as conn:
                customer_tables.ensure_table(conn, c.id, c.name)
    except Exception as e:
        print(f"Customer table create failed: {e}")
    return {"id": c.id, "name": c.name}


@router.patch("/customers/{customer_id}", dependencies=[Depends(require_staff), Depends(require_permission("edit"))])
def rename_customer(customer_id: int, data: schemas.CustomerRenameIn, db: Session = Depends(get_db),
                    user: models.User = Depends(get_current_user)):
    c = db.query(Customer).filter(Customer.id == customer_id, Customer.deleted_at.is_(None)).first()
    if not c:
        raise HTTPException(404, "Customer not found.")
    new_name = data.name.strip()
    if not new_name:
        raise HTTPException(400, "Customer name cannot be empty.")
    if len(new_name) > 150 or re.search(r"[;\`\'\"]", new_name):
        raise HTTPException(400, "Customer name contains invalid characters or exceeds 150 characters.")

    # Check if another active customer already uses this name
    existing = db.query(Customer).filter(
        Customer.id != customer_id,
        Customer.deleted_at.is_(None),
        or_(
            func.lower(Customer.name) == new_name.lower(),
            func.replace(func.lower(Customer.name), '_', ' ') == func.replace(new_name.lower(), '_', ' ')
        )
    ).first()
    if existing:
        raise HTTPException(400, f"A customer named '{existing.name}' already exists.")

    old_name = c.name
    c.name = new_name
    db.commit()

    # Synchronize dedicated table in cloud database
    try:
        customer_tables.rename_table(c.id, new_name, old_name)
    except Exception as e:
        print(f"Customer table rename error: {e}")

    log_activity(db, "customer_rename", f"Renamed customer ID {c.id} from '{old_name}' to '{new_name}'", user=user)
    db.commit()
    return {"id": c.id, "name": c.name, "old_name": old_name}


@router.post("/customers/rename-all", dependencies=[Depends(require_staff), Depends(require_permission("edit"))])
def rename_all_customers(data: schemas.CustomerRenameAllIn, db: Session = Depends(get_db),
                         user: models.User = Depends(get_current_user)):
    if not data.renames:
        raise HTTPException(400, "No customer renames provided.")

    seen_names = set()
    for item in data.renames:
        cid = item.id
        new_name = item.name.strip()
        if not new_name:
            raise HTTPException(400, f"Customer name for ID {cid} cannot be empty.")
        if len(new_name) > 150 or re.search(r"[;\`\'\"]", new_name):
            raise HTTPException(400, f"Customer name '{new_name}' contains invalid characters or exceeds 150 characters.")
        norm_name = new_name.lower().replace("_", " ")
        if norm_name in seen_names:
            raise HTTPException(400, f"Duplicate customer name '{new_name}' in request.")
        seen_names.add(norm_name)

    renamed_items = []
    for item in data.renames:
        cid = item.id
        new_name = item.name.strip()
        c = db.query(Customer).filter(Customer.id == cid, Customer.deleted_at.is_(None)).first()
        if not c:
            continue
        if c.name == new_name:
            continue
        old_name = c.name
        c.name = new_name
        db.flush()
        try:
            customer_tables.rename_table(c.id, new_name, old_name)
        except Exception as e:
            print(f"Error renaming table for {cid}: {e}")
        renamed_items.append({"id": cid, "old_name": old_name, "new_name": new_name})

    db.commit()
    log_activity(db, "customer_rename_all", f"Renamed {len(renamed_items)} customer portfolios", user=user)
    db.commit()
    return {"success": True, "renamed_count": len(renamed_items), "renames": renamed_items}


@router.post("/customers/{customer_id}/delete-all-entries", dependencies=[Depends(require_staff), Depends(require_permission("delete"))])
def delete_all_entries_for_customer(customer_id: int, data: schemas.CustomerDeleteAllEntriesIn, db: Session = Depends(get_db),
                                    user: models.User = Depends(get_current_user)):
    # Master Admin restriction: Master Admins are NOT permitted to delete valid inventory!
    if user.role == "master_admin" or not get_user_permissions(user).get("delete", False):
        raise HTTPException(status_code=403, detail="Master Admins are not permitted to delete valid inventory records.")

    c = db.query(Customer).filter(Customer.id == customer_id, Customer.deleted_at.is_(None)).first()
    if not c:
        raise HTTPException(404, "Customer not found.")

    clean_expected = c.name.strip().lower()
    clean_provided = data.confirm_name.strip().lower()
    if clean_provided != clean_expected and clean_provided.replace(" ", "_") != clean_expected.replace(" ", "_"):
        raise HTTPException(400, f"Confirmation name mismatch. Expected '{c.name}'.")

    total_to_delete = db.query(Device).filter(Device.customer_id == customer_id, ACTIVE).count()
    if total_to_delete == 0:
        return {"success": True, "customer_id": customer_id, "customer_name": c.name, "deleted_count": 0, "message": "No active entries found for this customer."}

    deleted_so_far = 0
    now_utc = datetime.utcnow()
    batch_size = 2000

    while True:
        batch_ids = [d_id for (d_id,) in db.query(Device.id).filter(Device.customer_id == customer_id, ACTIVE).limit(batch_size).all()]
        if not batch_ids:
            break
        if data.permanent:
            db.query(Device).filter(Device.id.in_(batch_ids)).delete(synchronize_session=False)
        else:
            db.query(Device).filter(Device.id.in_(batch_ids)).update({
                Device.deleted_at: now_utc,
                Device.deleted_by: user.user_id,
                Device.delete_reason: data.reason or f"Deleted all entries for customer '{c.name}'"
            }, synchronize_session=False)
        db.commit()
        deleted_so_far += len(batch_ids)

    try:
        customer_tables.clear_table(c.id, c.name)
    except Exception as e:
        print(f"Customer table clear rows error: {e}")

    log_activity(db, "customer_delete_all_entries", f"Deleted all {deleted_so_far} entries for customer '{c.name}' (ID {c.id})", user=user)
    db.commit()
    return {
        "success": True,
        "customer_id": c.id,
        "customer_name": c.name,
        "deleted_count": deleted_so_far
    }


@router.delete("/customers/{customer_id}", dependencies=[Depends(require_staff), Depends(require_permission("delete"))])
def delete_customer(customer_id: int, db: Session = Depends(get_db),
                    user: models.User = Depends(get_current_user)):
    if user.role == "master_admin" or not get_user_permissions(user).get("delete", False):
        raise HTTPException(status_code=403, detail="Master Admins are not permitted to delete customer portfolios.")

    c = db.query(Customer).filter(Customer.id == customer_id, Customer.deleted_at.is_(None)).first()
    if not c:
        raise HTTPException(404, "Customer not found.")

    c.deleted_at = datetime.utcnow()
    c.deleted_by = user.user_id

    db.query(Device).filter(Device.customer_id == customer_id, ACTIVE).update({
        Device.deleted_at: datetime.utcnow(),
        Device.deleted_by: user.user_id,
        Device.delete_reason: f"Customer '{c.name}' deleted"
    }, synchronize_session=False)

    log_activity(db, "customer_delete", f"Customer '{c.name}' (ID {c.id}) moved to bin", user=user)
    db.commit()
    return {"success": True, "customer_id": c.id, "customer_name": c.name}


@router.put("/customers/{customer_id}/avatar", dependencies=[Depends(require_staff), Depends(require_permission("edit"))])
@router.patch("/customers/{customer_id}/avatar", dependencies=[Depends(require_staff), Depends(require_permission("edit"))])
def update_customer_avatar(customer_id: int, data: dict, db: Session = Depends(get_db),
                           user: models.User = Depends(get_current_user)):
    c = db.query(Customer).filter(Customer.id == customer_id, Customer.deleted_at.is_(None)).first()
    if not c:
        raise HTTPException(404, "Customer not found.")
    c.avatar = data.get("avatar")
    db.commit()
    return {"id": c.id, "avatar": c.avatar}



# ---------- devices ----------
def build_device_query(db: Session, customer_id, q, column, sort_by, sort_dir, sim_type="all",
                       expiry_filter="all", expiry_days=30, exp_start=None, exp_end=None):
    """Search + sort shared by the table and the Excel export, so both always agree."""
    query = db.query(Device).join(Customer).options(joinedload(Device.customer)).filter(ACTIVE)
    if customer_id:
        query = query.filter(Device.customer_id == customer_id)

    if sim_type == "esim":
        query = query.filter(or_(
            func.upper(Device.sim_type) == 'E-SIM',
            and_(or_(Device.sim_type.is_(None), Device.sim_type == ''), Device.iccid.isnot(None), Device.iccid != '')
        ))
    elif sim_type == "plastic":
        query = query.filter(or_(
            func.upper(Device.sim_type).in_(['PLASTIC SIM', 'PLASTIC']),
            and_(or_(Device.sim_type.is_(None), Device.sim_type == ''), or_(Device.iccid.is_(None), Device.iccid == ''))
        ))

    today = date.today()
    if expiry_filter == "expired":
        query = query.filter(Device.expiry_date.isnot(None), Device.expiry_date < today)
    elif expiry_filter == "expiring_7":
        query = query.filter(Device.expiry_date.isnot(None), Device.expiry_date >= today, Device.expiry_date <= today + timedelta(days=7))
    elif expiry_filter == "expiring_15":
        query = query.filter(Device.expiry_date.isnot(None), Device.expiry_date >= today, Device.expiry_date <= today + timedelta(days=15))
    elif expiry_filter in ("expiring_30", "expiring_soon"):
        query = query.filter(Device.expiry_date.isnot(None), Device.expiry_date >= today, Device.expiry_date <= today + timedelta(days=expiry_days or 30))
    elif expiry_filter == "active":
        query = query.filter(Device.expiry_date.isnot(None), Device.expiry_date > today + timedelta(days=expiry_days or 30))
    elif expiry_filter == "unknown":
        query = query.filter(or_(Device.expiry_date.is_(None), Device.expiry_date == ""))
    elif expiry_filter == "custom" and exp_start and exp_end:
        query = query.filter(Device.expiry_date.isnot(None), Device.expiry_date >= exp_start, Device.expiry_date <= exp_end)

    q = (q or "").strip()
    if q:
        raw_parts = [p.strip() for p in re.split(r"[,;\r\n]+", q) if p.strip()]
        if not raw_parts:
            raw_parts = [q]

        col_map = {
            "imei": Device.imei,
            "iccid": Device.iccid,
            "ccid": Device.iccid,
            "device_sr_no": Device.device_sr_no,
            "msisdn1": Device.msisdn1,
            "msisdn2": Device.msisdn2,
            "customer_name": Customer.name,
            "model": Device.model,
            "sim_provider": Device.sim_provider,
            "activation_plan": Device.activation_plan,
            "firmware": Device.firmware,
            "flash": Device.flash,
            "sim_type": Device.sim_type,
        }

        term_filters = []
        for part in raw_parts:
            like = f"%{escape_like(part)}%"
            norm_dt = normalize_date_input(part)
            dt_like = f"%{norm_dt.isoformat()}%" if norm_dt else None
            clean_part = clean_iccid(part)

            if column and column.lower() in ("iccid", "ccid"):
                iccid_likes = [Device.iccid.like(like)]
                if clean_part:
                    iccid_likes.append(Device.iccid.like(f"%{escape_like(clean_part)}%"))
                    stripped = clean_part.rstrip("F")
                    if stripped and stripped != clean_part:
                        iccid_likes.append(Device.iccid.like(f"%{escape_like(stripped)}%"))
                    if len(clean_part) >= 19:
                        iccid_likes.append(func.left(Device.iccid, 19) == clean_part[:19])
                term_filters.append(or_(*iccid_likes))
            elif column and column in col_map:
                f_expr = col_map[column].like(like)
                if dt_like and column in ("act_date", "expiry_date", "dispatched_date"):
                    f_expr = or_(f_expr, col_map[column].like(dt_like))
                term_filters.append(f_expr)
            else:
                cols = [Device.device_sr_no, Device.imei, Device.iccid, Device.msisdn1, Device.msisdn2, Customer.name,
                        Device.model, Device.sim_provider, Device.activation_plan, Device.firmware, Device.flash, Device.sim_type]
                sub_or = [c.like(like) for c in cols] + [
                    cast(Device.act_date, String).like(like),
                    cast(Device.dispatched_date, String).like(like),
                    cast(Device.expiry_date, String).like(like),
                    cast(Device.extra, String).like(like)
                ]
                if clean_part:
                    sub_or.append(Device.iccid.like(f"%{escape_like(clean_part)}%"))
                    stripped = clean_part.rstrip("F")
                    if stripped and stripped != clean_part:
                        sub_or.append(Device.iccid.like(f"%{escape_like(stripped)}%"))
                if dt_like:
                    sub_or.extend([
                        cast(Device.act_date, String).like(dt_like),
                        cast(Device.dispatched_date, String).like(dt_like),
                        cast(Device.expiry_date, String).like(dt_like)
                    ])
                term_filters.append(or_(*sub_or))

        if len(term_filters) == 1:
            query = query.filter(term_filters[0])
        elif len(term_filters) > 1:
            query = query.filter(or_(*term_filters))

    sort_fields = {
        "id": Device.id,
        "imei": Device.imei,
        "customer_name": Customer.name,
        "device_sr_no": Device.device_sr_no,
        "iccid": Device.iccid,
        "ccid": Device.iccid,
        "act_date": Device.act_date,
        "expiry_date": Device.expiry_date,
        "model": Device.model,
        "sim_provider": Device.sim_provider,
    }
    col_attr = sort_fields.get((sort_by or "").lower(), Device.id)
    return query.order_by(desc(col_attr) if sort_dir == "desc" else asc(col_attr), asc(Device.id))


@router.get("/devices", dependencies=[Depends(require_permission("view"))])
def list_devices(customer_id: Optional[int] = None, q: str = "", column: Optional[str] = "all",
                 sim_type: Optional[str] = "all",
                 expiry_filter: Optional[str] = "all", expiry_days: int = Query(30, ge=1, le=365),
                 exp_start: Optional[date] = None, exp_end: Optional[date] = None,
                 sort_by: Optional[str] = "id", sort_dir: Optional[str] = "asc",
                 page: int = Query(1, ge=1), page_size: int = Query(100, ge=1, le=5000),
                 db: Session = Depends(get_db)):
    query = build_device_query(db, customer_id, q, column, sort_by, sort_dir, sim_type=sim_type,
                               expiry_filter=expiry_filter, expiry_days=expiry_days, exp_start=exp_start, exp_end=exp_end)
    total = query.count()
    items = query.offset((page - 1) * page_size).limit(page_size).all()
    return {"total": total, "page": page, "page_size": page_size,
            "pages": max(1, -(-total // page_size)), "items": [out(d) for d in items]}


@router.get("/devices/export", dependencies=[Depends(require_permission("download"))])
def export_devices(customer_id: Optional[int] = None, q: str = "", column: Optional[str] = "all",
                   sim_type: Optional[str] = "all",
                   expiry_filter: Optional[str] = "all", expiry_days: int = Query(30, ge=1, le=365),
                   exp_start: Optional[date] = None, exp_end: Optional[date] = None,
                   sort_by: Optional[str] = "id", sort_dir: Optional[str] = "asc",
                   title: str = "Customer_Inventory_Export",
                   db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """EVERY record that matches the current filters (not just one page) as an Excel file."""
    query = build_device_query(db, customer_id, q, column, sort_by, sort_dir, sim_type=sim_type,
                               expiry_filter=expiry_filter, expiry_days=expiry_days, exp_start=exp_start, exp_end=exp_end)
    extra_cols = [n for (n,) in db.query(models.ColumnRegistry.name).order_by(models.ColumnRegistry.name)]
    headers = ["ID"] + [label for _, label in EXPORT_COLS] + extra_cols
    rows, offset, batch = [], 0, 2000
    while True:
        part = query.offset(offset).limit(batch).all()
        if not part:
            break
        for d in part:
            o = out(d)
            rows.append([d.id] + [o.get(k) for k, _ in EXPORT_COLS] + [(d.extra or {}).get(c) for c in extra_cols])
        offset += batch
        for d in part:
            db.expunge(d)  # keep memory flat while reading tens of thousands of rows
    clean_title = re.sub(r"[^a-zA-Z0-9_-]", "_", title or "Export")
    filename = f"{clean_title}_{date.today().isoformat()}.xlsx"
    buf = make_excel(clean_title, headers, rows)
    log_activity(db, "export_excel", f"Exported {len(rows)} rows ({clean_title})", user=user)
    db.commit()
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/devices/expiry-summary", dependencies=[Depends(require_permission("view"))])
def get_expiry_summary(customer_id: Optional[int] = None, days: int = Query(30, ge=1, le=365), db: Session = Depends(get_db)):
    today = date.today()
    base_q = db.query(Device).filter(ACTIVE)
    if customer_id:
        base_q = base_q.filter(Device.customer_id == customer_id)

    total = base_q.count()
    expired = base_q.filter(Device.expiry_date.isnot(None), Device.expiry_date < today).count()
    exp_7 = base_q.filter(Device.expiry_date.isnot(None), Device.expiry_date >= today, Device.expiry_date <= today + timedelta(days=7)).count()
    exp_15 = base_q.filter(Device.expiry_date.isnot(None), Device.expiry_date >= today, Device.expiry_date <= today + timedelta(days=15)).count()
    exp_30 = base_q.filter(Device.expiry_date.isnot(None), Device.expiry_date >= today, Device.expiry_date <= today + timedelta(days=days)).count()
    active = base_q.filter(Device.expiry_date.isnot(None), Device.expiry_date > today + timedelta(days=days)).count()
    unknown = base_q.filter(or_(Device.expiry_date.is_(None), Device.expiry_date == "")).count()

    return {
        "total_count": total,
        "expired_count": expired,
        "expiring_7_days": exp_7,
        "expiring_15_days": exp_15,
        "expiring_30_days": exp_30,
        "active_count": active,
        "unknown_count": unknown,
        "alert_triggered": (expired > 0 or exp_30 > 0)
    }



@router.post("/devices/bulk/validate", dependencies=[Depends(require_permission("upload"))])
def bulk_validate(data: schemas.BulkIn, db: Session = Depends(get_db)):
    imeis = [(r.imei or "").strip() for r in data.rows if (r.imei or "").strip()]
    sr_nos = [(r.device_sr_no or "").strip() for r in data.rows if (r.device_sr_no or "").strip()]
    iccids = [(r.iccid or "").strip() for r in data.rows if (r.iccid or "").strip()]

    existing_devices_map = {}
    for i in range(0, len(imeis), 2000):
        for d in db.query(Device).filter(Device.imei.in_(imeis[i:i + 2000]), ACTIVE):
            existing_devices_map[d.imei] = d

    existing_sr_nos = {}
    for i in range(0, len(sr_nos), 2000):
        for (v, dev_imei) in db.query(Device.device_sr_no, Device.imei).filter(Device.device_sr_no.in_(sr_nos[i:i + 2000]), ACTIVE):
            if v:
                existing_sr_nos[v] = dev_imei

    existing_iccids = {}
    for i in range(0, len(iccids), 2000):
        for (v, dev_imei) in db.query(Device.iccid, Device.imei).filter(Device.iccid.in_(iccids[i:i + 2000]), ACTIVE):
            if v:
                existing_iccids[v] = dev_imei

    known_custs = {c.name.lower() for c in db.query(Customer.name).filter(Customer.deleted_at.is_(None))}
    known_custs_norm = {c.name.lower().replace('_', ' ') for c in db.query(Customer.name).filter(Customer.deleted_at.is_(None))}
    known_cust_map = {c.id: c.name for c in db.query(Customer.id, Customer.name).filter(Customer.deleted_at.is_(None))}
    known_cols = {n.lower() for (n,) in db.query(models.ColumnRegistry.name)}
    known_cols.update(c.lower() for c in BASE_COLUMNS)

    seen_imeis, seen_sr_nos, seen_iccids = set(), set(), set()
    new_entries, updated_entries, existing_entries = [], [], []
    valid, duplicates, invalids = [], [], []
    new_customers = set()
    new_columns = set()
    missing_columns = set()
    esim_count = 0
    plastic_count = 0

    has_imei_any = any(bool(r.imei) for r in data.rows)
    has_sr_any = any(bool(r.device_sr_no) for r in data.rows)
    if not has_imei_any:
        missing_columns.add("IMEI")
    if not has_sr_any and not getattr(data, "direct_empty", False):
        missing_columns.add("Device Sr No")

    for idx, r in enumerate(data.rows, start=1):
        errs = []
        cname = str(r.customer_name or "").strip()
        if not cname:
            errs.append("Missing or bad Customer Name")
        elif len(cname) > 100 or re.search(r"[;\`\'\"]", cname):
            errs.append("Customer name contains invalid characters or exceeds 100 characters")
        elif cname.lower() not in known_custs and cname.lower().replace('_', ' ') not in known_custs_norm:
            new_customers.add(cname)

        if not r.imei:
            errs.append("Missing IMEI")
        elif len(r.imei) > 255:
            errs.append("IMEI exceeds 255 characters")

        for fld, name in [
            (r.device_sr_no, "Device Sr No"),
            (r.iccid, "ICCID"),
            (r.msisdn1, "MSISDN1"),
            (r.msisdn2, "MSISDN2"),
            (r.model, "Model"),
            (r.sim_provider, "SIM Provider"),
            (r.activation_plan, "Activation Plan"),
            (r.firmware, "Firmware"),
            (r.flash, "Flash"),
        ]:
            if fld and len(str(fld)) > 255:
                errs.append(f"{name} exceeds 255 characters")

        if r.act_date and r.expiry_date and r.expiry_date < r.act_date:
            errs.append("Bad date sequence: Expiry date is before Activation date")

        # Check column names in extra
        if r.extra and isinstance(r.extra, dict):
            for k, v in r.extra.items():
                k_clean = str(k).strip()
                if not re.match(r"^[a-zA-Z0-9_\s-]+$", k_clean) or len(k_clean) > 64:
                    errs.append(f"Invalid column name '{k_clean}'")
                elif k_clean.lower() not in known_cols:
                    new_columns.add(k_clean)
                if v and len(str(v)) > 255:
                    errs.append(f"Value for '{k_clean}' exceeds 255 characters")

        if getattr(data, "direct_empty", False):
            stype = r.sim_type
        else:
            stype = r.sim_type or ("E-SIM" if r.iccid else "Plastic SIM")
        if stype == "E-SIM":
            esim_count += 1
        elif stype:
            plastic_count += 1

        row_dict = r.model_dump()
        row_dict["sim_type"] = stype
        row_dict["row_num"] = idx
        row_dict["expiry_status"] = compute_expiry_status(r.expiry_date, 30)

        mode = getattr(data, "mode", "skip") or "skip"

        # Check Duplicate Conflicts
        is_dup = False
        dup_reason = ""
        if r.imei and r.imei in seen_imeis:
            is_dup = True
            dup_reason = "Duplicate IMEI repeats earlier in this file"
        elif r.device_sr_no and r.device_sr_no in seen_sr_nos:
            is_dup = True
            dup_reason = "Duplicate Device Sr No repeats earlier in this file"
        elif r.device_sr_no and r.device_sr_no in existing_sr_nos and (existing_sr_nos.get(r.device_sr_no) != r.imei):
            is_dup = True
            dup_reason = f"Device Sr No '{r.device_sr_no}' already belongs to another IMEI in database"
        elif r.iccid and r.iccid in seen_iccids:
            is_dup = True
            dup_reason = "Duplicate ICCID repeats earlier in this file"
        elif r.iccid and r.iccid in existing_iccids and (existing_iccids.get(r.iccid) != r.imei):
            is_dup = True
            dup_reason = f"ICCID '{r.iccid}' already belongs to another IMEI in database"

        if is_dup:
            row_dict["entry_type"] = "Duplicate"
            row_dict["status"] = "Duplicate"
            row_dict["reason"] = dup_reason
            duplicates.append(row_dict)
        elif r.imei and r.imei in existing_devices_map:
            # Existing Device in database -> Check if updated or existing
            d = existing_devices_map[r.imei]
            differing = []
            for f in ["device_sr_no", "iccid", "msisdn1", "msisdn2", "model", "sim_provider", "activation_plan", "firmware", "flash", "sim_type"]:
                val = getattr(r, f, None)
                if val is not None and str(val).strip() != "":
                    old_val = getattr(d, f, None)
                    if str(val).strip() != str(old_val or "").strip():
                        differing.append(f)
            for f in ["act_date", "expiry_date", "dispatched_date"]:
                val = getattr(r, f, None)
                if val is not None:
                    old_val = getattr(d, f, None)
                    if val != old_val:
                        differing.append(f)
            if cname and d.customer_id in known_cust_map:
                existing_cname = known_cust_map[d.customer_id]
                if cname.lower().replace("_", " ") != existing_cname.lower().replace("_", " "):
                    differing.append("customer_name")

            if differing:
                row_dict["entry_type"] = "Updated"
                row_dict["status"] = "Updated"
                row_dict["reason"] = f"Existing in database — will update {len(differing)} field(s): {', '.join(differing)}"
                row_dict["differing_fields"] = differing
                updated_entries.append(row_dict)
                valid.append(row_dict)
            else:
                row_dict["entry_type"] = "Existing"
                row_dict["status"] = "Existing"
                row_dict["reason"] = "Already exists in database with identical values (No change needed)"
                existing_entries.append(row_dict)
                if mode == "replace":
                    valid.append(row_dict)
        else:
            # Completely New Entry!
            row_dict["entry_type"] = "New"
            row_dict["status"] = "New"
            row_dict["reason"] = "New entry — ready to add to database and customer table"
            new_entries.append(row_dict)
            valid.append(row_dict)

        if r.imei:
            seen_imeis.add(r.imei)
        if r.device_sr_no:
            seen_sr_nos.add(r.device_sr_no)
        if r.iccid:
            seen_iccids.add(r.iccid)

    all_rows = new_entries + updated_entries + existing_entries + duplicates + invalids
    all_rows.sort(key=lambda x: x.get("row_num", 0))

    cust_dups = {}
    for d in duplicates:
        cn = d.get("customer_name") or "General_Customer"
        cust_dups[cn] = cust_dups.get(cn, 0) + 1

    exp_expired_cnt = 0
    exp_soon_cnt = 0
    exp_active_cnt = 0
    exp_unknown_cnt = 0
    today = date.today()
    for r in data.rows:
        if not r.expiry_date:
            exp_unknown_cnt += 1
        elif r.expiry_date < today:
            exp_expired_cnt += 1
        elif r.expiry_date <= today + timedelta(days=30):
            exp_soon_cnt += 1
        else:
            exp_active_cnt += 1

    return {
        "valid_count": len(valid),
        "new_count": len(new_entries),
        "updated_count": len(updated_entries),
        "existing_count": len(existing_entries),
        "duplicate_count": len(duplicates),
        "invalid_count": len(invalids),
        "total_count": len(data.rows),
        "esim_count": esim_count,
        "plastic_count": plastic_count,
        "new_customers": sorted(list(new_customers)),
        "new_columns": sorted(list(new_columns)),
        "missing_columns_warning": list(missing_columns),
        "customer_duplicates": cust_dups,
        "expiry_counts": {
            "expired": exp_expired_cnt,
            "expiring_soon": exp_soon_cnt,
            "active": exp_active_cnt,
            "unknown": exp_unknown_cnt
        },
        "new_sample": new_entries[:300],
        "updated_sample": updated_entries[:300],
        "existing_sample": existing_entries[:300],
        "valid_sample": valid[:300],
        "duplicates": duplicates[:50000],
        "invalids": invalids[:2000],
        "all_sample": all_rows[:300],
        "new_imeis": [r.get("imei") for r in new_entries if r.get("imei")]
    }


def insert_chunk(db: Session, rows: list):
    """Inserts one chunk in a single multi-row INSERT and commits it. If an error occurs (such as
    invalid values, encoding or integrity issue in any row), falls back to row-by-row insertion
    so that only the bad record is skipped and every valid row in the chunk proceeds successfully."""
    if not rows:
        return
    for attempt in range(3):
        try:
            db.execute(insert(Device), rows)
            db.commit()
            return
        except OperationalError:
            db.rollback()
            if attempt == 2:
                raise
            time.sleep(2 * (attempt + 1))
        except Exception as e:
            # Multi-row batch failed; retry row-by-row to ignore invalid values and proceed
            db.rollback()
            for r in rows:
                try:
                    db.execute(insert(Device), [r])
                    db.commit()
                except Exception as inner_e:
                    db.rollback()
                    print(f"Skipping problematic row during bulk upload ({r.get('imei')}): {inner_e}")
            return


@router.post("/devices/bulk", dependencies=[Depends(require_permission("upload"))])
def bulk_add(data: schemas.BulkIn, db: Session = Depends(get_db),
             user: models.User = Depends(get_current_user)):
    imeis = [(r.imei or "").strip() for r in data.rows if (r.imei or "").strip()]
    sr_nos = [(r.device_sr_no or "").strip() for r in data.rows if (r.device_sr_no or "").strip()]
    iccids = [(r.iccid or "").strip() for r in data.rows if (r.iccid or "").strip()]
    mode = getattr(data, "mode", "skip") or "skip"

    existing_devices = {}
    for i in range(0, len(imeis), 2000):
        for d in db.query(Device).filter(Device.imei.in_(imeis[i:i + 2000])):
            existing_devices[d.imei] = d

    existing_sr_nos = {}
    for i in range(0, len(sr_nos), 2000):
        for (v, dev_imei) in db.query(models.Device.device_sr_no, models.Device.imei).filter(models.Device.device_sr_no.in_(sr_nos[i:i + 2000]), models.Device.deleted_at.is_(None)):
            if v:
                existing_sr_nos[v] = dev_imei

    existing_iccids = {}
    for i in range(0, len(iccids), 2000):
        for (v, dev_imei) in db.query(models.Device.iccid, models.Device.imei).filter(models.Device.iccid.in_(iccids[i:i + 2000]), models.Device.deleted_at.is_(None)):
            if v:
                existing_iccids[v] = dev_imei

    cache, added, updated, skipped, names = {}, 0, 0, [], set()
    seen_in_batch = set()
    seen_sr_nos = set()
    seen_iccids = set()
    mirror = []  # rows to copy into each customer's own table
    new_rows = []  # rows inserted in a few big batches instead of one round trip per row

    for r in data.rows:
        try:
            cust_name = str(r.customer_name or "").strip()
            if not cust_name:
                cust_name = "General_Customer"

            raw_imei = str(r.imei or "").strip()
            if "e+" in raw_imei.lower():
                try:
                    raw_imei = str(int(float(raw_imei)))
                except Exception:
                    pass
            imei = re.sub(r"[^0-9A-Za-z_-]", "", raw_imei)
            if not imei or len(imei) < 5:
                # Invalid rows are also inserted: assign unique fallback identifier so DB accepts it
                alt = re.sub(r"[^0-9A-Za-z_-]", "", str(r.device_sr_no or r.iccid or "").strip())
                if alt and len(alt) >= 5:
                    imei = alt[:20]
                else:
                    imei = f"INV_{uuid.uuid4().hex[:10].upper()}"
            imei = imei[:20]

            # Duplicates must NEVER be inserted
            if imei in seen_in_batch:
                skipped.append(imei)
                continue
            seen_in_batch.add(imei)

            sr_no = str(r.device_sr_no or "").strip()[:50]
            if sr_no:
                sr_owner = existing_sr_nos.get(sr_no)
                if sr_no in seen_sr_nos or (sr_owner and sr_owner != imei and mode == "skip"):
                    skipped.append(sr_no)
                    continue
                seen_sr_nos.add(sr_no)

            iccid_val = str(r.iccid or "").strip()[:30]
            if iccid_val:
                iccid_owner = existing_iccids.get(iccid_val)
                if iccid_val in seen_iccids or (iccid_owner and iccid_owner != imei and mode == "skip"):
                    skipped.append(iccid_val)
                    continue
                seen_iccids.add(iccid_val)

            key = cust_name.lower().replace("_", " ")
            if key not in cache:
                cache[key] = customer_by_name(db, cust_name)
            cust = cache[key]

            exp = r.expiry_date or (r.act_date + timedelta(days=365) if r.act_date else None)
            stype = (r.sim_type or classify_sim_type(r.iccid))[:20]
            extra = clean_extra(r.extra)
            names.update(extra)

            if imei in existing_devices:
                in_bin = getattr(existing_devices[imei], "deleted_at", None)  # deleted earlier: bring it back instead of skipping it
                if mode == "replace" or in_bin:
                    d = existing_devices[imei]
                    d.customer_id = cust.id
                    d.act_date = r.act_date or d.act_date
                    d.expiry_date = exp or d.expiry_date
                    d.dispatched_date = r.dispatched_date or d.dispatched_date
                    d.sim_type = stype or d.sim_type
                    for f in TEXT_FIELDS:
                        val = getattr(r, f, None)
                        if val is not None and str(val).strip() != "":
                            max_len = FIELD_MAX_LEN.get(f, 100)
                            setattr(d, f, str(val).strip()[:max_len])
                    if extra:
                        non_blank = {k: v for k, v in extra.items() if v is not None and str(v).strip() != ""}
                        merged = {**(d.extra or {}), **non_blank}
                        d.extra = merged or None
                    if d.deleted_at:
                        d.deleted_at = d.deleted_by = d.delete_reason = None
                    updated += 1
                    mirror.append((cust.id, cust.name, {
                        "imei": imei, "customer_name": cust.name, "act_date": d.act_date, "expiry_date": d.expiry_date, "dispatched_date": d.dispatched_date,
                        "sim_type": d.sim_type, "extra": d.extra or {}, **{f: getattr(d, f) for f in TEXT_FIELDS}}))
                else:
                    skipped.append(imei)
            else:
                row_vals = {}
                for f in TEXT_FIELDS:
                    val = getattr(r, f, None)
                    if val is not None and str(val).strip() != "":
                        max_len = FIELD_MAX_LEN.get(f, 100)
                        row_vals[f] = str(val).strip()[:max_len]
                    else:
                        row_vals[f] = None
                new_rows.append(dict(
                    customer_id=cust.id, imei=imei, act_date=r.act_date, expiry_date=exp,
                    dispatched_date=r.dispatched_date, sim_type=stype, extra=extra or None,
                    added_by=user.id, **row_vals
                ))
                existing_devices[imei] = True
                added += 1
                mirror.append((cust.id, cust.name, {
                    "imei": imei, "customer_name": cust.name, "act_date": r.act_date, "expiry_date": exp, "dispatched_date": r.dispatched_date,
                    "sim_type": stype, "extra": extra, **row_vals}))
        except Exception as row_err:
            print(f"Skipping row error in bulk_add: {row_err}")
            continue

    try:
        db.commit()  # customers first
    except Exception as ce:
        db.rollback()
        print(f"Commit error during bulk_add customers: {ce}")

    saved = 0
    try:
        for i in range(0, len(new_rows), 500):  # pymysql turns each list into ONE multi-row INSERT
            chunk = new_rows[i:i + 500]
            insert_chunk(db, chunk)
            saved += len(chunk)
        register_columns(db, names)
        if len(data.rows) == 1:
            log_activity(db, "device_add", f"IMEI {data.rows[0].imei}" + (" (updated)" if updated else "" if added else " (skipped)"), user=user)
        else:
            log_activity(db, "device_upload", f"{saved} added, {updated} updated, {len(skipped)} duplicates skipped (batch of {len(data.rows)})", user=user)
        db.commit()

        # Clean up any staged error records for successfully added or updated devices
        if saved > 0 or updated > 0:
            success_imeis = [r["imei"] for r in new_rows] + list(existing_devices.keys())
            for i in range(0, len(success_imeis), 1000):
                db.query(models.ErroredRecord).filter(models.ErroredRecord.imei.in_(success_imeis[i:i + 1000])).delete(synchronize_session=False)
            db.commit()
    except OperationalError:
        db.rollback()
        raise HTTPException(503, f"The database did not answer in time. {saved} of {len(new_rows)} new records were saved. "
                                 "Upload the same file again: records already saved are skipped automatically. "
                                 "Then run  python customer_tables.py  once to fill the customer tables.")

    try:  # per-customer tables; a failure here never breaks the upload itself
        db.commit()
        if getattr(getattr(db, "bind", None), "dialect", None) and getattr(db.bind.dialect, "name", "") == "sqlite":
            customer_tables.mirror_rows(mirror, conn=db.connection())
            db.commit()
        else:
            customer_tables.mirror_rows(mirror)
    except Exception as e:
        print(f"Customer table sync failed: {e}")
    return {"added": added, "updated": updated, "skipped_duplicates": len(skipped), "skipped_imeis": skipped[:50]}


@router.post("/devices/sheet-upload", dependencies=[Depends(require_permission("upload"))])
def sheet_upload(data: schemas.BulkIn, db: Session = Depends(get_db),
                 user: models.User = Depends(get_current_user)):
    return bulk_add(data, db, user)


@router.post("/devices", status_code=201, dependencies=[Depends(require_permission("add"))])
def add_device(data: schemas.DeviceIn, db: Session = Depends(get_db),
               user: models.User = Depends(get_current_user)):
    if not data.imei or len(data.imei) < 5 or not data.customer_name:
        raise HTTPException(400, "Customer name and a valid IMEI are required.")
    res = bulk_add(schemas.BulkIn(rows=[data]), db, user)
    if not res["added"]:
        raise HTTPException(400, "A device with this IMEI already exists (it may be in Deleted Customers).")
    return res


def clean_imei(val) -> str:
    if not val:
        return ""
    s = str(val).strip()
    if re.search(r"e\+", s, re.I):
        try:
            return str(int(float(s)))
        except Exception:
            pass
    s = re.sub(r"[^0-9A-Za-z_-]", "", s)
    return s.strip()


def find_devices_for_sim_data(db: Session, file_sims: dict, match_by: str) -> dict[str, list[models.Device]]:
    """Fetches and maps devices for SIM data upload.
    When matching by ICCID:
    Matches against all records where ICCID matches the uploaded file's ICCID:
    1. Exact match (case-insensitive)
    2. Stripped of trailing 'F' / 'f'
    3. Standard 19-digit core
    Returns a dict mapping file_key -> list of matching Device objects.
    """
    devs_by_key = {}
    if not file_sims:
        return devs_by_key

    if match_by == "imei":
        all_keys = list(file_sims.keys())
        matched_devices = []
        chunk_size = 1000
        for i in range(0, len(all_keys), chunk_size):
            sub_keys = all_keys[i:i + chunk_size]
            found = db.query(Device).options(joinedload(Device.customer)).filter(
                ACTIVE,
                Device.imei.in_(sub_keys)
            ).all()
            matched_devices.extend(found)
        for d in matched_devices:
            if not d.imei:
                continue
            k = clean_imei(d.imei)
            devs_by_key.setdefault(k, []).append(d)
        return devs_by_key

    # ICCID Matching
    search_keys_set = set()
    file_keys_exact = {}
    file_keys_rstrip = {}
    file_keys_prefix19 = {}

    for fk in file_sims.keys():
        cleaned_fk = clean_iccid(fk)
        file_keys_exact.setdefault(cleaned_fk, []).append(fk)
        search_keys_set.add(cleaned_fk)

        stripped = cleaned_fk.rstrip("F")
        file_keys_rstrip.setdefault(stripped, []).append(fk)
        search_keys_set.add(stripped)
        search_keys_set.add(stripped + "F")

        if len(cleaned_fk) >= 19:
            p19 = cleaned_fk[:19]
            file_keys_prefix19.setdefault(p19, []).append(fk)
            search_keys_set.add(p19)
            search_keys_set.add(p19 + "F")

    all_search_keys = list(search_keys_set)
    matched_devices = []
    chunk_size = 1000
    for i in range(0, len(all_search_keys), chunk_size):
        sub = all_search_keys[i:i + chunk_size]
        sub_19 = [k[:19] for k in sub if len(k) >= 19]
        filters = [Device.iccid.in_(sub)]
        if sub_19:
            filters.append(func.left(Device.iccid, 19).in_(sub_19))
        found = db.query(Device).options(joinedload(Device.customer)).filter(
            ACTIVE,
            or_(*filters)
        ).all()
        matched_devices.extend(found)

    seen_ids = set()
    unique_devices = []
    for d in matched_devices:
        if d.id not in seen_ids:
            seen_ids.add(d.id)
            unique_devices.append(d)

    for d in unique_devices:
        if not d.iccid:
            continue
        dicc = clean_iccid(d.iccid)
        target_fks = file_keys_exact.get(dicc)
        if not target_fks:
            target_fks = file_keys_rstrip.get(dicc.rstrip("F"))
        if not target_fks and len(dicc) >= 19:
            target_fks = file_keys_prefix19.get(dicc[:19])

        if target_fks:
            for fk in set(target_fks):
                devs_by_key.setdefault(fk, []).append(d)

    return devs_by_key


@router.post("/simdata/check", dependencies=[Depends(require_permission("search"))])
def check_sim_data(data: schemas.SimDataCheckIn, db: Session = Depends(get_db),
                    user: models.User = Depends(get_current_user)):
    rows = data.rows
    overwrite = data.overwrite
    chosen_provider = (data.sim_provider or "").strip()
    match_by = "iccid"  # SIM Data upload strictly matches by ICCID; IMEI option disabled and removed

    # Verify that either a provider, ICCID, or at least one data column is provided
    has_any_data = any(
        (r.iccid or r.msisdn1 or r.msisdn2 or r.act_date or r.expiry_date or r.model or r.firmware or r.flash or r.activation_plan or r.sim_type or r.sim_provider or r.extra)
        for r in rows
    )
    if not chosen_provider and not has_any_data:
        raise HTTPException(400, "Refused: The file contains no ICCID, MSISDN, date, or SIM data columns and no SIM Provider was selected.")

    total_file_rows = len(rows)
    skipped_rows = 0
    file_sims = {}
    repeated_count = 0
    unreadable_dates = 0

    for r in rows:
        if match_by == "imei":
            key_val = clean_imei(r.imei)
            raw_key = key_val
            if not key_val or len(key_val) < 5:
                skipped_rows += 1
                continue
            match_key = key_val
        else:
            raw_iccid = clean_iccid(r.iccid)
            if not raw_iccid or len(raw_iccid) < 5:
                skipped_rows += 1
                continue
            raw_key = raw_iccid
            match_key = raw_iccid

        row_dict = {
            "raw_key": raw_key,
            "raw_imei": clean_imei(r.imei),
            "raw_iccid": clean_iccid(r.iccid),
            "msisdn1": clean_str(r.msisdn1),
            "msisdn2": clean_str(r.msisdn2),
            "act_date": r.act_date,
            "expiry_date": r.expiry_date,
            "sim_provider": chosen_provider or clean_str(r.sim_provider),
            "extra": r.extra or {}
        }

        if match_key in file_sims:
            repeated_count += 1
            existing = file_sims[match_key]
            for f in ("msisdn1", "msisdn2", "act_date", "expiry_date", "sim_provider", "raw_iccid", "raw_imei"):
                if not existing.get(f) and row_dict.get(f):
                    existing[f] = row_dict[f]
            if row_dict.get("extra"):
                existing["extra"] = {**(existing.get("extra") or {}), **row_dict["extra"]}
        else:
            file_sims[match_key] = row_dict

    devs_by_key = find_devices_for_sim_data(db, file_sims, match_by)

    keys_found_set = set()
    records_matched = 0
    records_to_update = 0
    records_already_same = 0
    sample_changes = []
    detailed_comparison = []
    cust_stats = {}

    for key, frow in file_sims.items():
        dev_list = devs_by_key.get(key, [])
        if not dev_list:
            continue

        keys_found_set.add(frow["raw_key"])
        records_matched += len(dev_list)

        for d in dev_list:
            cid = d.customer_id
            cname = d.customer.name if d.customer else "Unknown"
            if cid not in cust_stats:
                cust_stats[cid] = {
                    "customer_id": cid,
                    "customer_name": cname,
                    "matched_count": 0,
                    "to_update_count": 0,
                    "already_same_count": 0,
                    "msisdn1_changes": 0,
                    "msisdn2_changes": 0,
                    "act_date_changes": 0,
                    "expiry_date_changes": 0,
                    "sim_provider_changes": 0,
                    "iccid_changes": 0
                }
            cs = cust_stats[cid]
            cs["matched_count"] += 1

            has_change = False
            changed_fields_list = []
            changed_labels_list = []

            # Compare activation date, expiry date, msisdn1, msisdn2, sim_provider
            fields_to_check = [
                ("act_date", "ACT DATE"),
                ("expiry_date", "EXP DATE"),
                ("msisdn1", "MSISDN1"),
                ("msisdn2", "MSISDN2"),
            ]
            if chosen_provider or frow.get("sim_provider"):
                fields_to_check.append(("sim_provider", "SIM PROVIDER"))
            if match_by == "imei" and frow.get("raw_iccid"):
                fields_to_check.append(("iccid", "ICCID"))

            field_diff_details = {}

            for field_name, label in fields_to_check:
                if field_name == "iccid":
                    file_val = frow.get("raw_iccid")
                else:
                    file_val = frow.get(field_name)

                curr_val = getattr(d, field_name, None)
                curr_str = curr_val.isoformat() if isinstance(curr_val, date) else str(curr_val or "").strip()
                new_str = file_val.isoformat() if isinstance(file_val, date) else str(file_val or "").strip()

                if file_val is None or file_val == "":
                    # Empty in file -> keep existing unchanged
                    field_diff_details[field_name] = {
                        "old": curr_str or "—",
                        "new": curr_str or "—",
                        "changed": False
                    }
                    continue

                is_changed = False
                if overwrite:
                    if curr_str != new_str:
                        is_changed = True
                else:
                    if not curr_str and new_str:
                        is_changed = True

                field_diff_details[field_name] = {
                    "old": curr_str or "—",
                    "new": new_str or "—",
                    "changed": is_changed
                }

                if is_changed:
                    has_change = True
                    changed_fields_list.append(field_name)
                    changed_labels_list.append(label)
                    stat_key = f"{field_name}_changes"
                    if stat_key in cs:
                        cs[stat_key] += 1
                    if len(sample_changes) < 300:
                        sample_changes.append({
                            "customer_name": cname,
                            "imei": d.imei,
                            "iccid": d.iccid,
                            "field": label,
                            "old_val": curr_str or "—",
                            "new_val": new_str
                        })

            if has_change:
                records_to_update += 1
                cs["to_update_count"] += 1
            else:
                records_already_same += 1
                cs["already_same_count"] += 1

            if len(detailed_comparison) < 500:
                detailed_comparison.append({
                    "imei": d.imei or "—",
                    "iccid": d.iccid or "—",
                    "customer_name": cname,
                    "status": "different" if has_change else "same",
                    "changed_fields": changed_fields_list,
                    "changed_labels": changed_labels_list,
                    "act_date": field_diff_details.get("act_date", {"old": "—", "new": "—", "changed": False}),
                    "expiry_date": field_diff_details.get("expiry_date", {"old": "—", "new": "—", "changed": False}),
                    "msisdn1": field_diff_details.get("msisdn1", {"old": "—", "new": "—", "changed": False}),
                    "msisdn2": field_diff_details.get("msisdn2", {"old": "—", "new": "—", "changed": False}),
                    "sim_provider": field_diff_details.get("sim_provider", {"old": d.sim_provider or "—", "new": frow.get("sim_provider") or d.sim_provider or "—", "changed": field_diff_details.get("sim_provider", {}).get("changed", False)})
                })

    not_found_keys = [frow["raw_key"] for key, frow in file_sims.items() if frow["raw_key"] not in keys_found_set]
    columns_to_create = ["MSISDN1", "MSISDN2", "ACT DATE", "EXP DATE", "SIM PROVIDER"]

    not_found_records = []
    for key, frow in file_sims.items():
        if frow["raw_key"] not in keys_found_set:
            act_d = frow.get("act_date")
            exp_d = frow.get("expiry_date")
            not_found_records.append({
                "iccid": frow["raw_key"],
                "msisdn1": frow.get("msisdn1") or "—",
                "msisdn2": frow.get("msisdn2") or "—",
                "act_date": act_d.isoformat() if isinstance(act_d, date) else (str(act_d) if act_d else "—"),
                "expiry_date": exp_d.isoformat() if isinstance(exp_d, date) else (str(exp_d) if exp_d else "—"),
                "sim_provider": frow.get("sim_provider") or "—",
                "status": "Not Found in Any Customer (Will NOT be uploaded)"
            })

    return {
        "file_rows_count": total_file_rows,
        "match_by": match_by,
        "keys_found_count": len(keys_found_set),
        "iccids_found": len(keys_found_set),
        "records_matched": records_matched,
        "records_to_update": records_to_update,
        "records_already_same": records_already_same,
        "not_found_keys_count": len(not_found_keys),
        "iccids_not_found_count": len(not_found_keys),
        "not_found_keys": not_found_keys,
        "iccids_not_found": not_found_keys,
        "not_found_records": not_found_records,
        "skipped_rows": skipped_rows,
        "repeated_keys_count": repeated_count,
        "repeated_iccids_count": repeated_count,
        "can_insert_count": 0,
        "unreadable_dates_count": unreadable_dates,
        "columns_to_create": columns_to_create,
        "customer_breakdown": sorted(list(cust_stats.values()), key=lambda x: x["customer_name"]),
        "sample_changes": sample_changes,
        "detailed_comparison": detailed_comparison
    }


@router.post("/simdata/apply", dependencies=[Depends(require_permission("edit"))])
def apply_sim_data(data: schemas.SimDataApplyIn, db: Session = Depends(get_db),
                   user: models.User = Depends(get_current_user)):
    rows = data.rows
    overwrite = data.overwrite
    chosen_provider = (data.sim_provider or "").strip()
    match_by = "iccid"  # SIM Data upload strictly matches by ICCID

    has_any_data = any(
        (r.iccid or r.msisdn1 or r.msisdn2 or r.act_date or r.expiry_date or r.sim_provider or r.imei or r.extra)
        for r in rows
    )
    if not chosen_provider and not has_any_data and not getattr(data, "insert_new", False):
        raise HTTPException(400, "Refused: No ICCID, data columns, or SIM Provider specified.")

    file_sims = {}
    for r in rows:
        if match_by == "imei":
            key_val = clean_imei(r.imei)
            raw_key = key_val
            if not key_val or len(key_val) < 5:
                continue
            match_key = key_val
        else:
            raw_iccid = clean_iccid(r.iccid)
            if not raw_iccid or len(raw_iccid) < 5:
                continue
            raw_key = raw_iccid
            match_key = raw_iccid

        row_dict = {
            "raw_key": raw_key,
            "raw_imei": clean_imei(r.imei),
            "raw_iccid": clean_iccid(r.iccid),
            "msisdn1": clean_str(r.msisdn1),
            "msisdn2": clean_str(r.msisdn2),
            "act_date": r.act_date,
            "expiry_date": r.expiry_date,
            "sim_provider": chosen_provider or clean_str(r.sim_provider),
            "extra": r.extra or {}
        }
        if match_key in file_sims:
            existing = file_sims[match_key]
            for f in ("msisdn1", "msisdn2", "act_date", "expiry_date", "sim_provider", "raw_iccid", "raw_imei"):
                if not existing.get(f) and row_dict.get(f):
                    existing[f] = row_dict[f]
            if row_dict.get("extra"):
                existing["extra"] = {**(existing.get("extra") or {}), **row_dict["extra"]}
        else:
            file_sims[match_key] = row_dict

    devs_by_key = find_devices_for_sim_data(db, file_sims, match_by)

    updated_count = 0
    already_same_count = 0
    matched_keys = set()
    mirror_items = []
    applied_comparison = []

    for key, frow in file_sims.items():
        dev_list = devs_by_key.get(key, [])
        if not dev_list:
            continue

        matched_keys.add(frow["raw_key"])

        for d in dev_list:
            changed = False
            field_diff_details = {}
            changed_fields_list = []

            for field in ("msisdn1", "msisdn2", "act_date", "expiry_date", "sim_provider"):
                val = frow.get(field)
                curr = getattr(d, field)
                curr_cmp = curr.isoformat() if isinstance(curr, date) else str(curr or "").strip()
                val_cmp = val.isoformat() if isinstance(val, date) else str(val or "").strip()

                if val is None or val == "":
                    field_diff_details[field] = {
                        "old": curr_cmp or "—",
                        "new": curr_cmp or "—",
                        "changed": False
                    }
                    continue

                is_field_changed = False
                if overwrite:
                    if curr_cmp != val_cmp:
                        setattr(d, field, val)
                        is_field_changed = True
                        changed = True
                else:
                    if not curr_cmp and val_cmp:
                        setattr(d, field, val)
                        is_field_changed = True
                        changed = True

                field_diff_details[field] = {
                    "old": curr_cmp or "—",
                    "new": val_cmp or "—",
                    "changed": is_field_changed
                }
                if is_field_changed:
                    changed_fields_list.append(field)

            if match_by == "iccid" and overwrite and frow.get("raw_iccid"):
                curr_iccid = clean_iccid(d.iccid)
                new_iccid = frow["raw_iccid"]
                if curr_iccid != new_iccid and len(new_iccid) >= len(curr_iccid):
                    d.iccid = new_iccid
                    changed = True

            cust_name = d.customer.name if d.customer else "Unknown"
            if changed:
                updated_count += 1
                cust_id = d.customer_id
                row_mirror = {
                    "imei": d.imei,
                    "customer_name": cust_name,
                    "device_sr_no": d.device_sr_no,
                    "iccid": d.iccid,
                    "msisdn1": d.msisdn1,
                    "msisdn2": d.msisdn2,
                    "act_date": d.act_date,
                    "expiry_date": d.expiry_date,
                    "dispatched_date": d.dispatched_date,
                    "model": d.model,
                    "sim_provider": d.sim_provider,
                    "activation_plan": d.activation_plan,
                    "firmware": d.firmware,
                    "flash": d.flash,
                    "sim_type": d.sim_type,
                    "extra": d.extra or {}
                }
                mirror_items.append((cust_id, cust_name, row_mirror))
            else:
                already_same_count += 1

            if len(applied_comparison) < 1500:
                applied_comparison.append({
                    "imei": d.imei or "—",
                    "iccid": d.iccid or "—",
                    "customer_name": cust_name,
                    "status": "updated" if changed else "same",
                    "changed_fields": changed_fields_list,
                    "act_date": field_diff_details.get("act_date", {"old": "—", "new": "—", "changed": False}),
                    "expiry_date": field_diff_details.get("expiry_date", {"old": "—", "new": "—", "changed": False}),
                    "msisdn1": field_diff_details.get("msisdn1", {"old": "—", "new": "—", "changed": False}),
                    "msisdn2": field_diff_details.get("msisdn2", {"old": "—", "new": "—", "changed": False}),
                    "sim_provider": field_diff_details.get("sim_provider", {"old": d.sim_provider or "—", "new": d.sim_provider or "—", "changed": field_diff_details.get("sim_provider", {}).get("changed", False)})
                })

    unmatched_keys = []
    not_found_records = []
    for key, frow in file_sims.items():
        if frow["raw_key"] not in matched_keys:
            unmatched_keys.append(frow["raw_key"])
            act_d = frow.get("act_date")
            exp_d = frow.get("expiry_date")
            not_found_records.append({
                "iccid": frow["raw_key"],
                "msisdn1": frow.get("msisdn1") or "—",
                "msisdn2": frow.get("msisdn2") or "—",
                "act_date": act_d.isoformat() if isinstance(act_d, date) else (str(act_d) if act_d else "—"),
                "expiry_date": exp_d.isoformat() if isinstance(exp_d, date) else (str(exp_d) if exp_d else "—"),
                "sim_provider": frow.get("sim_provider") or "—",
                "status": "Not Found in Any Customer (Skipped / Not Uploaded)"
            })

    # STRICT REQUIREMENT: If ICCID is not found, DO NOT upload/insert into any customer!
    # inserted_count is strictly 0. Only existing customer devices matched by ICCID are updated.

    db.commit()

    if mirror_items:
        try:
            if getattr(getattr(db, "bind", None), "dialect", None) and getattr(db.bind.dialect, "name", "") == "sqlite":
                customer_tables.mirror_rows(mirror_items, conn=db.connection())
                db.commit()
            else:
                customer_tables.mirror_rows(mirror_items)
        except Exception as me:
            print(f"Mirroring SIM updates to customer tables warning: {me}")

    log_msg = f"Updated {updated_count} devices matched by ICCID from SIM provider {chosen_provider or 'file'}. Ignored {len(unmatched_keys)} ICCIDs not found in any customer."
    log_activity(db, "simdata_apply", log_msg, user=user)
    db.commit()

    return {
        "updated": updated_count,
        "inserted": 0,
        "already_same": already_same_count,
        "ignored_identical": already_same_count,
        "match_by": match_by,
        "unmatched_count": len(unmatched_keys),
        "unmatched_keys": unmatched_keys,
        "unmatched_iccids": unmatched_keys,
        "not_found_count": len(unmatched_keys),
        "not_found_records": not_found_records,
        "applied_comparison": applied_comparison
    }


@router.post("/fetch/bulk", dependencies=[Depends(require_permission("search"))])
def bulk_fetch(data: schemas.BulkFetchIn, db: Session = Depends(get_db)):
    raw_terms = [q.strip() for q in data.queries if q.strip()]
    if not raw_terms:
        return {
            "exact_matches": [], "similar_matches": [],
            "total_matches": 0, "values_searched_count": 0,
            "values_found_count": 0, "not_found_values": []
        }

    seen = set()
    terms = []
    for t in raw_terms:
        if t not in seen:
            seen.add(t)
            terms.append(t)

    field = (data.field or "all").strip().lower()
    if field in ("iccid", "ccid"):
        field = "iccid"
    cust_id = data.customer_id
    sim_type = data.sim_type or "all"

    base_query = db.query(Device).options(joinedload(Device.customer)).filter(ACTIVE)
    if cust_id:
        base_query = base_query.filter(Device.customer_id == cust_id)

    if sim_type == "esim":
        base_query = base_query.filter(or_(
            func.upper(Device.sim_type) == 'E-SIM',
            and_(or_(Device.sim_type.is_(None), Device.sim_type == ''), Device.iccid.isnot(None), Device.iccid != '')
        ))
    elif sim_type == "plastic":
        base_query = base_query.filter(or_(
            func.upper(Device.sim_type).in_(['PLASTIC SIM', 'PLASTIC']),
            and_(or_(Device.sim_type.is_(None), Device.sim_type == ''), or_(Device.iccid.is_(None), Device.iccid == ''))
        ))

    exact_items = []
    chunk_size = 2000
    for i in range(0, len(terms), chunk_size):
        sub_terms = terms[i:i + chunk_size]
        q = base_query
        if field == "imei":
            q = q.filter(Device.imei.in_(sub_terms))
        elif field in ("iccid", "ccid"):
            clean_terms = [clean_iccid(x) for x in sub_terms if x]
            stripped_terms = [c.rstrip("F") for c in clean_terms if c]
            f_terms = [c.rstrip("F") + "F" for c in clean_terms if c]
            p19 = [c[:19] for c in clean_terms if len(c) >= 19]
            iccid_search = list(set(sub_terms + clean_terms + stripped_terms + f_terms))
            iccid_conds = [Device.iccid.in_(iccid_search)]
            if p19:
                iccid_conds.append(func.left(Device.iccid, 19).in_(p19))
            q = q.filter(or_(*iccid_conds))
        elif field == "device_sr_no":
            q = q.filter(Device.device_sr_no.in_(sub_terms))
        elif field == "msisdn1":
            q = q.filter(Device.msisdn1.in_(sub_terms))
        elif field == "customer_name":
            q = q.join(Customer).filter(Customer.name.in_(sub_terms))
        else:
            clean_terms = [clean_iccid(x) for x in sub_terms if x]
            stripped_terms = [c.rstrip("F") for c in clean_terms if c]
            f_terms = [c.rstrip("F") + "F" for c in clean_terms if c]
            iccid_search = list(set(sub_terms + clean_terms + stripped_terms + f_terms))
            q = q.filter(or_(
                Device.imei.in_(sub_terms),
                Device.iccid.in_(iccid_search),
                Device.device_sr_no.in_(sub_terms),
                Device.msisdn1.in_(sub_terms),
                Device.msisdn2.in_(sub_terms)
            ))
        exact_items.extend(q.all())

    exact_ids = {d.id for d in exact_items}

    matched_term_set = set()
    for d in exact_items:
        d_iccid_clean = clean_iccid(d.iccid) if d.iccid else ""
        for t in (d.imei, d.iccid, d.device_sr_no, d.msisdn1, d.msisdn2, (d.customer.name if d.customer else "")):
            if t and t in seen:
                matched_term_set.add(t)
            elif t and clean_iccid(t) in seen:
                matched_term_set.add(clean_iccid(t))
        for t in terms:
            t_clean = clean_iccid(t)
            if t_clean and d_iccid_clean and (
                t_clean == d_iccid_clean or
                t_clean.rstrip("F") == d_iccid_clean.rstrip("F") or
                (len(t_clean) >= 19 and len(d_iccid_clean) >= 19 and t_clean[:19] == d_iccid_clean[:19])
            ):
                matched_term_set.add(t)

    unmatched_terms = [t for t in terms if t not in matched_term_set]
    similar_items = []
    similar_terms_to_check = unmatched_terms[:300]

    for t in similar_terms_to_check:
        like = f"%{escape_like(t)}%"
        sim_q = base_query
        if field == "imei":
            sim_q = sim_q.filter(Device.imei.like(like))
        elif field in ("iccid", "ccid"):
            clean_t = clean_iccid(t)
            sim_conds = [Device.iccid.like(like)]
            if clean_t:
                sim_conds.append(Device.iccid.like(f"%{escape_like(clean_t)}%"))
                if clean_t.rstrip("F") != clean_t:
                    sim_conds.append(Device.iccid.like(f"%{escape_like(clean_t.rstrip('F'))}%"))
            sim_q = sim_q.filter(or_(*sim_conds))
        elif field == "customer_name":
            sim_q = sim_q.join(Customer).filter(Customer.name.like(like))
        else:
            clean_t = clean_iccid(t)
            sim_or = [
                Device.imei.like(like),
                Device.iccid.like(like),
                Device.device_sr_no.like(like),
                Device.msisdn1.like(like)
            ]
            if clean_t:
                sim_or.append(Device.iccid.like(f"%{escape_like(clean_t)}%"))
            sim_q = sim_q.filter(or_(*sim_or))
        found = sim_q.limit(10).all()
        for s in found:
            if s.id not in exact_ids and s.id not in {x.id for x in similar_items}:
                similar_items.append(s)
                matched_term_set.add(t)

    not_found_values = [t for t in terms if t not in matched_term_set]

    all_returned = exact_items + similar_items
    all_imeis = [d.imei for d in all_returned if d.imei]
    all_iccids = [d.iccid for d in all_returned if d.iccid]
    all_srnos = [d.device_sr_no for d in all_returned if d.device_sr_no]

    dup_imeis = {i for i, in db.query(Device.imei).filter(ACTIVE, Device.imei.in_(all_imeis)).group_by(Device.imei).having(func.count(Device.id) > 1)} if all_imeis else set()
    dup_iccids = {i for i, in db.query(Device.iccid).filter(ACTIVE, Device.iccid.in_(all_iccids)).group_by(Device.iccid).having(func.count(Device.id) > 1)} if all_iccids else set()
    dup_srnos = {i for i, in db.query(Device.device_sr_no).filter(ACTIVE, Device.device_sr_no.in_(all_srnos)).group_by(Device.device_sr_no).having(func.count(Device.id) > 1)} if all_srnos else set()

    def format_item(d, is_exact):
        res = out(d)
        res["is_exact"] = is_exact
        res["has_duplicates"] = (d.imei in dup_imeis or d.iccid in dup_iccids or d.device_sr_no in dup_srnos)
        res["duplicate_fields"] = [f for f, is_dup in [("IMEI", d.imei in dup_imeis), ("ICCID", d.iccid in dup_iccids), ("Sr No", d.device_sr_no in dup_srnos)] if is_dup]
        return res

    exact_res = [format_item(d, True) for d in exact_items]
    similar_res = [format_item(d, False) for d in similar_items]

    return {
        "exact_matches": exact_res,
        "similar_matches": similar_res,
        "total_matches": len(exact_res) + len(similar_res),
        "values_searched_count": len(terms),
        "values_found_count": len(terms) - len(not_found_values),
        "not_found_values": not_found_values
    }


@router.get("/devices/quick-search", dependencies=[Depends(require_permission("search"))])
def quick_search(q: str = Query("", min_length=1), db: Session = Depends(get_db)):
    term = q.strip()
    if not term:
        return {"customers": [], "devices": []}
    like = f"%{escape_like(term)}%"

    custs = db.query(Customer).filter(Customer.deleted_at.is_(None), Customer.name.like(like)).limit(5).all()

    devs = (db.query(Device).options(joinedload(Device.customer))
            .filter(ACTIVE, or_(
                Device.imei == term,
                Device.iccid == term,
                Device.device_sr_no == term,
                Device.imei.like(like),
                Device.iccid.like(like),
                Device.device_sr_no.like(like),
                Device.msisdn1.like(like),
                Device.model.like(like)
            )).limit(10).all())

    return {
        "customers": [{"id": c.id, "name": c.name, "exact": c.name.lower() == term.lower()} for c in custs],
        "devices": [{**out(d), "exact": term.lower() in [str(d.imei or '').lower(), str(d.iccid or '').lower(), str(d.device_sr_no or '').lower()]} for d in devs]
    }


@router.get("/inventory/insights", dependencies=[Depends(require_permission("view"))])
def inventory_insights(db: Session = Depends(get_db)):
    today = date.today()
    soon_30 = today + timedelta(days=30)
    over_4m = today + timedelta(days=120)

    total_devices = db.query(func.count(Device.id)).filter(ACTIVE).scalar() or 0
    total_customers = db.query(func.count(Customer.id)).filter(Customer.deleted_at.is_(None)).scalar() or 0
    expired = db.query(func.count(Device.id)).filter(ACTIVE, Device.expiry_date < today).scalar() or 0
    expiring_30 = db.query(func.count(Device.id)).filter(ACTIVE, Device.expiry_date >= today, Device.expiry_date <= soon_30).scalar() or 0
    expiring_over_4m = db.query(func.count(Device.id)).filter(ACTIVE, Device.expiry_date > over_4m).scalar() or 0

    timeline_raw = (db.query(func.date_format(Device.expiry_date, '%Y-%m').label('month'), func.count(Device.id))
                    .filter(ACTIVE, Device.expiry_date.isnot(None))
                    .group_by('month').order_by('month').all())
    timeline = [{"month": m, "count": c} for m, c in timeline_raw if m]

    by_cust = (db.query(Customer.name, func.count(Device.id).label('cnt'))
               .join(Device, (Device.customer_id == Customer.id) & ACTIVE)
               .filter(Customer.deleted_at.is_(None))
               .group_by(Customer.id, Customer.name).order_by(desc('cnt')).limit(10).all())

    by_provider = (db.query(func.coalesce(Device.sim_provider, 'Not specified'), func.count(Device.id))
                   .filter(ACTIVE).group_by(Device.sim_provider).order_by(desc(func.count(Device.id))).all())

    by_model = (db.query(func.coalesce(Device.model, 'Not specified'), func.count(Device.id))
                .filter(ACTIVE).group_by(Device.model).order_by(desc(func.count(Device.id))).all())

    by_plan = (db.query(func.coalesce(Device.activation_plan, 'Not specified'), func.count(Device.id))
               .filter(ACTIVE).group_by(Device.activation_plan).order_by(desc(func.count(Device.id))).all())

    recent = (db.query(Device).options(joinedload(Device.customer))
              .filter(ACTIVE).order_by(Device.id.desc()).limit(10).all())

    user_ids = {d.added_by for d in recent if d.added_by}
    user_map = {u.id: u.name for u in db.query(models.User).filter(models.User.id.in_(list(user_ids))).all()} if user_ids else {}

    recent_list = []
    for d in recent:
        item = out(d)
        item["added_by_name"] = user_map.get(d.added_by, f"User #{d.added_by}" if d.added_by else "System / Bulk")
        recent_list.append(item)

    return {
        "total_devices": total_devices,
        "total_customers": total_customers,
        "expired": expired,
        "expiring_30": expiring_30,
        "expiring_over_4m": expiring_over_4m,
        "timeline": timeline,
        "by_customer": [{"name": n, "count": c} for n, c in by_cust],
        "by_provider": [{"name": n, "count": c} for n, c in by_provider[:10]],
        "by_model": [{"name": n, "count": c} for n, c in by_model[:10]],
        "by_plan": [{"name": n, "count": c} for n, c in by_plan[:10]],
        "recent_added": recent_list
    }


@router.get("/inventory/health", dependencies=[Depends(require_permission("view"))])
def inventory_health(db: Session = Depends(get_db)):
    def find_dups(column):
        return db.query(column, func.count(Device.id)).filter(ACTIVE, column.isnot(None), column != '').group_by(column).having(func.count(Device.id) > 1).all()

    dup_imei = find_dups(Device.imei)
    dup_iccid = find_dups(Device.iccid)
    dup_m1 = find_dups(Device.msisdn1)
    dup_m2 = find_dups(Device.msisdn2)
    dup_sr = find_dups(Device.device_sr_no)

    total = db.query(func.count(Device.id)).filter(ACTIVE).scalar() or 0

    bad_imeis = (db.query(Device).options(joinedload(Device.customer))
                 .filter(ACTIVE, or_(func.length(Device.imei) != 15, Device.imei.op('REGEXP')('^[0-9]+$') == 0))
                 .limit(100).all())

    bad_dates = (db.query(Device).options(joinedload(Device.customer))
                 .filter(ACTIVE, Device.act_date.isnot(None), Device.expiry_date.isnot(None), Device.expiry_date < Device.act_date)
                 .limit(100).all())

    cols = ["customer_name", "device_sr_no", "imei", "iccid", "msisdn1", "msisdn2", "act_date", "expiry_date",
            "dispatched_date", "model", "sim_provider", "activation_plan", "firmware", "flash", "sim_type"]
    completeness = {}
    if total > 0:
        for c in cols:
            if c == "customer_name":
                cnt = db.query(func.count(Device.id)).join(Customer).filter(ACTIVE, Customer.name.isnot(None), Customer.name != '').scalar() or 0
            else:
                attr = getattr(Device, c)
                cnt = db.query(func.count(Device.id)).filter(ACTIVE, attr.isnot(None), attr != '').scalar() or 0
            completeness[c] = round((cnt / total) * 100, 1)

    issues = []
    for val, count in dup_imei[:20]:
        issues.append({"type": "Duplicate IMEI", "identifier": val, "count": count, "description": f"IMEI {val} appears {count} times"})
    for val, count in dup_iccid[:20]:
        issues.append({"type": "Duplicate ICCID", "identifier": val, "count": count, "description": f"ICCID {val} appears {count} times"})
    for val, count in dup_sr[:20]:
        issues.append({"type": "Duplicate Sr No", "identifier": val, "count": count, "description": f"Device Sr No {val} appears {count} times"})
    for d in bad_imeis[:20]:
        issues.append({"type": "Invalid IMEI Length", "identifier": d.imei, "count": 1, "description": f"Device #{d.id} IMEI '{d.imei}' has length {len(d.imei or '')} (expected 15 digits)"})
    for d in bad_dates[:20]:
        issues.append({"type": "Expiry Before Activation", "identifier": d.imei, "count": 1, "description": f"Device #{d.id} expires on {d.expiry_date} before activation on {d.act_date}"})

    return {
        "total_records": total,
        "duplicate_imei_count": len(dup_imei),
        "duplicate_iccid_count": len(dup_iccid),
        "duplicate_msisdn1_count": len(dup_m1),
        "duplicate_msisdn2_count": len(dup_m2),
        "duplicate_sr_count": len(dup_sr),
        "bad_imei_count": len(bad_imeis),
        "bad_date_count": len(bad_dates),
        "completeness": completeness,
        "issues": issues
    }


@router.get("/inventory/report", dependencies=[Depends(require_permission("download"))])
def inventory_report(db: Session = Depends(get_db)):
    today = date.today()
    soon_30 = today + timedelta(days=30)

    rows = (db.query(
                Customer.id, Customer.name,
                func.count(Device.id).label('total'),
                func.sum(case((Device.expiry_date > soon_30, 1), else_=0)).label('active'),
                func.sum(case(((Device.expiry_date >= today) & (Device.expiry_date <= soon_30), 1), else_=0)).label('expiring'),
                func.sum(case((Device.expiry_date < today, 1), else_=0)).label('expired')
            )
            .outerjoin(Device, (Device.customer_id == Customer.id) & ACTIVE)
            .filter(Customer.deleted_at.is_(None))
            .group_by(Customer.id, Customer.name).order_by(Customer.name).all())

    return [{"id": i, "name": n, "total": int(t or 0), "active": int(a or 0), "expiring_30": int(e or 0), "expired": int(x or 0)} for i, n, t, a, e, x in rows]


@router.post("/export", dependencies=[Depends(require_permission("download"))])
def export_excel(data: schemas.ExportIn, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    buf = make_excel(data.title or "CustomerVault_Export", data.headers, data.rows)
    clean_title = re.sub(r'[^a-zA-Z0-9_-]', '_', data.title or 'Export')
    filename = f"{clean_title}_{date.today().isoformat()}.xlsx"
    log_activity(db, "export_excel", f"Exported {len(data.rows)} rows ({data.title})", user=user)
    db.commit()
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{filename}\""}
    )


@router.post("/columns/add", dependencies=[Depends(require_permission("add"))])
def add_column(data: schemas.AddColumnIn, db: Session = Depends(get_db),
               user: models.User = Depends(get_current_user)):
    col_name = data.column_name.strip()
    if not col_name:
        raise HTTPException(400, "Column name cannot be blank.")
    norm = col_name.lower().replace(' ', '_')
    if norm in BASE_COLUMNS:
        raise HTTPException(400, f"'{col_name}' is a core system column and cannot be added as a custom column.")

    # 1. Register the column in ColumnRegistry
    register_columns(db, [col_name])

    # 2. Select target devices
    query = db.query(Device).options(joinedload(Device.customer)).filter(ACTIVE)
    if data.target == "customer" and data.customer_id:
        query = query.filter(Device.customer_id == data.customer_id)
    elif data.target == "specified":
        raw_imeis = data.imeis or []
        cleaned_imeis = [re.sub(r"\D", "", str(s)) for s in raw_imeis if re.sub(r"\D", "", str(s))]
        if not cleaned_imeis:
            raise HTTPException(400, "Please provide at least one valid IMEI for specified entries.")
        query = query.filter(Device.imei.in_(cleaned_imeis))
        if data.customer_id:
            query = query.filter(Device.customer_id == data.customer_id)
    elif data.target == "all" and data.customer_id:
        query = query.filter(Device.customer_id == data.customer_id)

    val_to_set = str(data.value).strip() if data.value is not None else ""
    if data.value is None or (data.target != "specified" and not val_to_set):
        db.commit()
        log_activity(db, "column_add", f"Added column '{col_name}'", user=user)
        return {"column": col_name, "value": None, "records_affected": 0}

    # 2. Fast SQL JSON_SET update
    json_path = f"$.\"{col_name}\""
    where_clauses = ["deleted_at IS NULL"]
    params = {"path": json_path, "val": val_to_set}

    if data.target == "customer" and data.customer_id:
        where_clauses.append("customer_id = :cid")
        params["cid"] = data.customer_id
    elif data.target == "specified":
        raw_imeis = data.imeis or []
        cleaned_imeis = [re.sub(r"\D", "", str(s)) for s in raw_imeis if re.sub(r"\D", "", str(s))]
        if not cleaned_imeis:
            raise HTTPException(400, "Please provide at least one valid IMEI for specified entries.")
        where_clauses.append("imei IN :imeis")
        params["imeis"] = tuple(cleaned_imeis)
        if data.customer_id:
            where_clauses.append("customer_id = :cid")
            params["cid"] = data.customer_id
    elif data.target == "all" and data.customer_id:
        where_clauses.append("customer_id = :cid")
        params["cid"] = data.customer_id

    sql = f"UPDATE sim_devices SET extra = JSON_SET(COALESCE(extra, '{{}}'), :path, :val) WHERE {' AND '.join(where_clauses)}"
    res = db.execute(text(sql), params)
    count_affected = res.rowcount
    db.commit()

    # 3. Mirror to customer tables
    if data.target == "customer" and data.customer_id:
        cust = db.query(models.Customer).filter(models.Customer.id == data.customer_id).first()
        if cust:
            def _mirror_cust(cid, cname, col, v):
                try:
                    with customer_tables.engine.begin() as conn:
                        t = customer_tables.ensure_table(conn, cid, cname, [col])
                        c_col = customer_tables.extra_col(col)
                        conn.execute(text(f"UPDATE `{t}` SET `{c_col}` = :v"), {"v": v})
                except Exception as e:
                    print(f"Error updating customer table for column add: {e}")
            import threading
            threading.Thread(target=_mirror_cust, args=(cust.id, cust.name, col_name, val_to_set), daemon=True).start()
    elif data.target == "specified":
        cleaned_imeis = [re.sub(r"\D", "", str(s)) for s in (data.imeis or []) if re.sub(r"\D", "", str(s))]
        if cleaned_imeis:
            def _mirror_specified(imeis_list):
                try:
                    with SessionLocal() as s_db:
                        devs = s_db.query(Device).options(joinedload(Device.customer)).filter(Device.imei.in_(imeis_list), ACTIVE).all()
                        m_items = []
                        for d in devs:
                            cname = d.customer.name if d.customer else "Unknown"
                            r_dict = {
                                "imei": d.imei,
                                "device_sr_no": d.device_sr_no,
                                "iccid": d.iccid,
                                "msisdn1": d.msisdn1,
                                "msisdn2": d.msisdn2,
                                "act_date": d.act_date,
                                "expiry_date": d.expiry_date,
                                "dispatched_date": d.dispatched_date,
                                "model": d.model,
                                "sim_provider": d.sim_provider,
                                "activation_plan": d.activation_plan,
                                "firmware": d.firmware,
                                "flash": d.flash,
                                "sim_type": d.sim_type,
                                "extra": dict(d.extra or {})
                            }
                            m_items.append((d.customer_id, cname, r_dict))
                        if m_items:
                            customer_tables.mirror_rows(m_items)
                except Exception as e:
                    print(f"Error mirroring specified items: {e}")
            import threading
            threading.Thread(target=_mirror_specified, args=(cleaned_imeis,), daemon=True).start()

    log_activity(db, "column_add", f"Added column '{col_name}' (value '{val_to_set}' applied to {count_affected} records)", user=user)
    return {"column": col_name, "value": val_to_set, "records_affected": count_affected}


@router.post("/columns/delete", dependencies=[Depends(require_admin), Depends(require_permission("delete"))])
def delete_column(data: schemas.DeleteColumnIn, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    col_name = data.column_name.strip()
    norm = col_name.lower().replace(' ', '_')
    if norm in BASE_COLUMNS:
        raise HTTPException(400, f"Cannot delete '{col_name}' because it is a core system column.")

    entry = db.query(models.ColumnRegistry).filter(models.ColumnRegistry.name == col_name).first()
    if entry:
        db.delete(entry)

    json_path = f"$.\"{col_name}\""
    sql = "UPDATE sim_devices SET extra = JSON_REMOVE(extra, :path) WHERE extra IS NOT NULL AND JSON_CONTAINS_PATH(extra, 'one', :path)"
    res = db.execute(text(sql), {"path": json_path})
    count_cleaned = res.rowcount

    log_activity(db, "column_delete", f"Deleted column '{col_name}' (cleared from {count_cleaned} records)", user=admin)
    db.commit()
    return {"deleted": col_name, "records_affected": count_cleaned}


@router.get("/devices/lookup", dependencies=[Depends(require_permission("view"))])
def lookup(value: str, db: Session = Depends(get_db)):
    v = value.strip()
    items = (db.query(Device).options(joinedload(Device.customer))
             .filter(ACTIVE, or_(Device.imei == v, Device.iccid == v, Device.msisdn1 == v,
                                 Device.msisdn2 == v, Device.device_sr_no == v)).limit(50).all())
    return [out(d) for d in items]


@router.get("/devices/{device_id}", dependencies=[Depends(require_permission("view"))])
def get_device_by_id(device_id: int, db: Session = Depends(get_db)):
    d = db.query(Device).options(joinedload(Device.customer)).filter(Device.id == device_id, ACTIVE).first()
    if not d:
        raise HTTPException(404, "Device not found.")
    return out(d)


@router.put("/devices/{device_id}", dependencies=[Depends(require_permission("edit"))])
def update_device_by_id(device_id: int, data: schemas.DeviceUpdate, db: Session = Depends(get_db),
                        user: models.User = Depends(get_current_user)):
    d = db.query(Device).options(joinedload(Device.customer)).filter(Device.id == device_id, ACTIVE).first()
    if not d:
        raise HTTPException(404, "Device not found.")
    changes = data.model_dump(exclude_unset=True)
    fields = sorted(changes)
    if "extra" in changes:
        merged = merge_extra(d.extra, changes.pop("extra"))
        d.extra = merged or None
        register_columns(db, merged)
    if "customer_name" in changes:
        cn = (changes.pop("customer_name") or "").strip()
        if cn:
            d.customer_id = customer_by_name(db, cn).id
    for k, v in changes.items():
        setattr(d, k, v)
    log_activity(db, "device_edit", f"Device #{d.id} (IMEI {d.imei}): {', '.join(fields) or 'no changes'}", user=user)
    db.commit()
    db.refresh(d)
    try:
        cname = d.customer.name if d.customer else "customer"
        row = {"imei": d.imei, "extra": d.extra or {}}
        row.update({f: getattr(d, f) for f in customer_tables.BASE_NAMES})
        customer_tables.mirror_single_device(d.customer_id, cname, row)
    except Exception as e:
        print(f"Customer table mirror failed: {e}")
    return out(d)


@router.get("/devices/by-imei/{imei}", dependencies=[Depends(require_permission("view"))])
def get_by_imei(imei: str, db: Session = Depends(get_db)):
    return out(active_by_imei(db, imei))


@router.put("/devices/by-imei/{imei}", dependencies=[Depends(require_permission("edit"))])
def update_by_imei(imei: str, data: schemas.DeviceUpdate, db: Session = Depends(get_db),
                   user: models.User = Depends(get_current_user)):
    d = active_by_imei(db, imei)
    changes = data.model_dump(exclude_unset=True)
    fields = sorted(changes)
    if "extra" in changes:
        merged = merge_extra(d.extra, changes.pop("extra"))
        d.extra = merged or None
        register_columns(db, merged)
    if "customer_name" in changes:
        cn = (changes.pop("customer_name") or "").strip()
        if cn:
            d.customer_id = customer_by_name(db, cn).id
    for k, v in changes.items():
        setattr(d, k, v)
    log_activity(db, "device_edit", f"IMEI {d.imei}: {', '.join(fields) or 'no changes'}", user=user)
    db.commit()
    db.refresh(d)
    try:
        cname = d.customer.name if d.customer else "customer"
        row = {"imei": d.imei, "extra": d.extra or {}}
        row.update({f: getattr(d, f) for f in customer_tables.BASE_NAMES})
        customer_tables.mirror_single_device(d.customer_id, cname, row)
    except Exception as e:
        print(f"Customer table mirror failed: {e}")
    return out(d)


@router.delete("/devices/by-imei/{imei}", dependencies=[Depends(require_staff), Depends(require_permission("delete"))])
def delete_by_imei(imei: str, data: Optional[schemas.ReasonIn] = None, db: Session = Depends(get_db),
                   user: models.User = Depends(get_current_user)):
    d = active_by_imei(db, imei)
    cid = d.customer_id
    trash(d, user, (data.reason if data else None) or "Deleted by IMEI")
    log_activity(db, "device_delete", f"IMEI {d.imei} moved to bin", user=user)
    db.commit()
    try:
        customer_tables.delete_device_from_table(cid, d.imei)
    except Exception as e:
        print(f"Customer table delete device failed: {e}")
    return {"moved_to_bin": d.imei}


@router.delete("/devices/{device_id}", dependencies=[Depends(require_staff), Depends(require_permission("delete"))])
def delete_device(device_id: int, db: Session = Depends(get_db),
                  user: models.User = Depends(get_current_user)):
    d = db.query(Device).filter(Device.id == device_id, ACTIVE).first()
    if not d:
        raise HTTPException(404, "Record not found.")
    cid = d.customer_id
    imei = d.imei
    trash(d, user, "Deleted from list")
    log_activity(db, "device_delete", f"IMEI {d.imei} moved to bin", user=user)
    db.commit()
    try:
        customer_tables.delete_device_from_table(cid, imei)
    except Exception as e:
        print(f"Customer table delete device failed: {e}")
    return {"moved_to_bin": d.id}


# ---------- expiry alerts ----------
@router.get("/alerts/expiry", dependencies=[Depends(require_permission("view"))])
def expiry_alerts(days: int = Query(30, ge=0, le=365), customer_id: Optional[int] = None, limit: int = 5000, db: Session = Depends(get_db)):
    q = (db.query(Device).options(joinedload(Device.customer))
         .filter(ACTIVE, Device.expiry_date.isnot(None), Device.expiry_date <= date.today() + timedelta(days=days)))
    if customer_id:
        q = q.filter(Device.customer_id == customer_id)
    items = q.order_by(Device.expiry_date).limit(limit).all()
    return [out(d) for d in items]


# ---------- deleted customers (recycle bin) ----------
@router.get("/bin", dependencies=[Depends(require_staff), Depends(require_permission("view"))])
def bin_list(db: Session = Depends(get_db)):
    items = (db.query(Device).options(joinedload(Device.customer))
             .filter(Device.deleted_at.isnot(None)).order_by(Device.deleted_at.desc()).all())
    return [out(d) for d in items]


@router.post("/bin/{device_id}/restore", dependencies=[Depends(require_staff), Depends(require_permission("edit"))])
def bin_restore(device_id: int, db: Session = Depends(get_db),
                user: models.User = Depends(get_current_user)):
    d = db.query(Device).filter(Device.id == device_id, Device.deleted_at.isnot(None)).first()
    if not d:
        raise HTTPException(404, "Not in the bin.")
    d.deleted_at = d.deleted_by = d.delete_reason = None
    if d.customer and d.customer.deleted_at:
        d.customer.deleted_at = d.customer.deleted_by = None
    log_activity(db, "device_restore", f"IMEI {d.imei} restored", user=user)
    db.commit()
    try:
        cname = d.customer.name if d.customer else ""
        customer_tables.mirror_single_device(d.customer_id, cname, out(d))
    except Exception as e:
        print(f"Customer table mirror single device on restore failed: {e}")
    return {"restored": d.id}


@router.delete("/bin/{device_id}", dependencies=[Depends(require_admin), Depends(require_permission("delete"))])
def bin_purge(device_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    d = db.query(Device).filter(Device.id == device_id, Device.deleted_at.isnot(None)).first()
    if not d:
        raise HTTPException(404, "Not in the bin.")
    cid, imei = d.customer_id, d.imei
    cname = d.customer.name if d.customer else ""
    log_activity(db, "device_purge", f"IMEI {d.imei} deleted forever", user=admin)
    db.delete(d)
    db.commit()
    try:
        customer_tables.delete_device_from_table(cid, imei, cname)
    except Exception as e:
        print(f"Customer table delete device failed: {e}")
    return {"deleted_forever": device_id}


@router.delete("/bin", dependencies=[Depends(require_admin), Depends(require_permission("delete"))])
def bin_empty(limit: int = 0, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    total_remaining_devices = db.query(Device).filter(Device.deleted_at.isnot(None)).count()
    deleted_custs = db.query(Customer).filter(Customer.deleted_at.isnot(None)).all()

    if limit > 0:
        dev_objs = db.query(Device).options(joinedload(Device.customer)).filter(Device.deleted_at.isnot(None)).limit(limit).all()
        dev_ids = [d.id for d in dev_objs]
        for d in dev_objs:
            try:
                cname = d.customer.name if d.customer else ""
                customer_tables.delete_device_from_table(d.customer_id, d.imei, cname)
            except Exception:
                pass
        if dev_ids:
            n = db.query(Device).filter(Device.id.in_(dev_ids)).delete(synchronize_session=False)
        else:
            n = 0
        db.commit()
        remaining = max(0, total_remaining_devices - n)
        cust_purged = 0
        if remaining == 0 and deleted_custs:
            for c in deleted_custs:
                try:
                    customer_tables.drop_table(c.id, c.name)
                except Exception as e:
                    print(f"Drop table on empty cart failed: {e}")
                db.delete(c)
            cust_purged = len(deleted_custs)
            db.commit()
        log_activity(db, "bin_empty_chunk", f"Permanently deleted {n} records from bin ({remaining} remaining)", user=admin)
        return {"deleted_forever": n, "remaining_count": remaining, "customers_purged": cust_purged}
    else:
        # Drop tables for deleted customers permanently
        for c in deleted_custs:
            try:
                customer_tables.drop_table(c.id, c.name)
            except Exception as e:
                print(f"Drop table on empty cart failed: {e}")
            db.delete(c)
        dev_objs = db.query(Device).options(joinedload(Device.customer)).filter(Device.deleted_at.isnot(None)).all()
        for d in dev_objs:
            try:
                cname = d.customer.name if d.customer else ""
                customer_tables.delete_device_from_table(d.customer_id, d.imei, cname)
            except Exception:
                pass
        n = db.query(Device).filter(Device.deleted_at.isnot(None)).delete(synchronize_session=False)
        log_activity(db, "bin_empty", f"Emptied the bin ({n} records and {len(deleted_custs)} customers deleted forever)", user=admin)
        db.commit()
        return {"deleted_forever": n, "remaining_count": 0, "customers_purged": len(deleted_custs)}


# ---------- errored / invalid records staging ----------
def errored_out(r: models.ErroredRecord):
    return {
        "id": r.id,
        "customer_name": r.customer_name or "",
        "imei": r.imei or "",
        "device_sr_no": r.device_sr_no or "",
        "iccid": r.iccid or "",
        "msisdn1": r.msisdn1 or "",
        "msisdn2": r.msisdn2 or "",
        "act_date": r.act_date.isoformat() if r.act_date else None,
        "expiry_date": r.expiry_date.isoformat() if r.expiry_date else None,
        "dispatched_date": r.dispatched_date.isoformat() if r.dispatched_date else None,
        "model": r.model or "",
        "sim_provider": r.sim_provider or "",
        "activation_plan": r.activation_plan or "",
        "firmware": r.firmware or "",
        "flash": r.flash or "",
        "sim_type": r.sim_type or "",
        "extra": r.extra or {},
        "error_reason": r.error_reason or "Invalid data",
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


@router.get("/errors", dependencies=[Depends(require_permission("view"))])
def list_errors(reason: Optional[str] = None, q: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(models.ErroredRecord)
    if reason and reason.strip():
        query = query.filter(models.ErroredRecord.error_reason.ilike(f"%{reason.strip()}%"))
    if q and q.strip():
        term = f"%{q.strip()}%"
        query = query.filter(
            or_(
                models.ErroredRecord.imei.ilike(term),
                models.ErroredRecord.customer_name.ilike(term),
                models.ErroredRecord.iccid.ilike(term),
                models.ErroredRecord.device_sr_no.ilike(term),
                models.ErroredRecord.msisdn1.ilike(term),
                models.ErroredRecord.error_reason.ilike(term)
            )
        )
    rows = query.order_by(models.ErroredRecord.created_at.desc(), models.ErroredRecord.id.desc()).all()
    return [errored_out(r) for r in rows]


@router.get("/errors/summary", dependencies=[Depends(require_permission("view"))])
def errors_summary(db: Session = Depends(get_db)):
    rows = db.query(models.ErroredRecord.error_reason, func.count(models.ErroredRecord.id)).group_by(models.ErroredRecord.error_reason).all()
    total = sum(c for _, c in rows)
    reasons = [{"reason": r or "Unknown Error", "count": c} for r, c in rows]
    reasons.sort(key=lambda x: x["count"], reverse=True)
    return {"total": total, "reasons": reasons}


@router.post("/errors", dependencies=[Depends(require_permission("view"))])
def save_errors(data: schemas.BulkErrorsIn, db: Session = Depends(get_db)):
    added = 0
    for item in data.records:
        rec = models.ErroredRecord(
            customer_name=item.customer_name,
            imei=item.imei,
            device_sr_no=item.device_sr_no,
            iccid=item.iccid,
            msisdn1=item.msisdn1,
            msisdn2=item.msisdn2,
            act_date=item.act_date,
            expiry_date=item.expiry_date,
            dispatched_date=item.dispatched_date,
            model=item.model,
            sim_provider=item.sim_provider,
            activation_plan=item.activation_plan,
            firmware=item.firmware,
            flash=item.flash,
            sim_type=item.sim_type,
            extra=item.extra,
            error_reason=item.error_reason or "Invalid entry"
        )
        db.add(rec)
        added += 1
    db.commit()
    return {"saved_errors": added}


@router.put("/errors/{error_id}/upload", dependencies=[Depends(require_permission("upload"))])
def upload_corrected_error(error_id: int, data: schemas.DeviceIn, db: Session = Depends(get_db),
                           user: models.User = Depends(get_current_user)):
    err = db.query(models.ErroredRecord).filter(models.ErroredRecord.id == error_id).first()
    if not err:
        raise HTTPException(404, "Errored record not found.")

    cname = (data.customer_name or err.customer_name or "").strip() or "General_Customer"
    imei_val = (data.imei or "").strip()
    if not imei_val or len(imei_val) < 5:
        alt = clean_imei(data.device_sr_no or data.iccid or err.device_sr_no or err.iccid or "")
        if alt and len(alt) >= 5:
            imei_val = alt[:20]
        else:
            imei_val = f"INV_{uuid.uuid4().hex[:10].upper()}"
    data.imei = imei_val
    data.customer_name = cname

    # Ingest using bulk_add logic so customer and TiDB Cloud table are synchronized
    res = bulk_add(schemas.BulkIn(rows=[data], mode="replace"), db=db, user=user)

    # Remove from errored records table
    db.delete(err)
    db.commit()

    log_activity(db, "error_resolved", f"Corrected and uploaded record IMEI {data.imei} to customer {data.customer_name}", user=user)
    return {"success": True, "result": res}


@router.post("/errors/{error_id}/push", dependencies=[Depends(require_permission("upload"))])
def push_error_to_customer(error_id: int, data: Optional[schemas.PushErrorIn] = None,
                           db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Pushes an errored entry directly into the specified customer even if it contains errors or duplicate attributes."""
    err = db.query(models.ErroredRecord).filter(models.ErroredRecord.id == error_id).first()
    if not err:
        raise HTTPException(404, "Errored record not found.")

    target_cname = ""
    if data and data.customer_name:
        target_cname = data.customer_name.strip()
    if not target_cname:
        target_cname = (err.customer_name or "").strip()
    if not target_cname:
        target_cname = "General_Customer"

    cust = customer_by_name(db, target_cname)

    raw_imei = clean_imei(err.imei) if err.imei else ""
    extra = dict(err.extra or {})
    extra["pushed_with_error"] = True
    if err.error_reason:
        extra["original_error"] = err.error_reason

    if not raw_imei or len(raw_imei) < 5:
        alt = clean_imei(err.device_sr_no or err.iccid or "")
        if alt and len(alt) >= 5:
            raw_imei = alt[:20]
        else:
            raw_imei = f"INV_{uuid.uuid4().hex[:10].upper()}"

    final_imei = raw_imei[:20]
    stype = err.sim_type or ("E-SIM" if err.iccid else "Plastic SIM")
    existing_dev = db.query(Device).filter(Device.imei == final_imei).first()
    if existing_dev:
        dev = existing_dev
        dev.customer_id = cust.id
        if err.device_sr_no:
            dev.device_sr_no = err.device_sr_no
        if err.iccid:
            dev.iccid = err.iccid
        if err.msisdn1:
            dev.msisdn1 = err.msisdn1
        if err.msisdn2:
            dev.msisdn2 = err.msisdn2
        if err.act_date:
            dev.act_date = err.act_date
        if err.expiry_date:
            dev.expiry_date = err.expiry_date
        if err.dispatched_date:
            dev.dispatched_date = err.dispatched_date
        if err.model:
            dev.model = err.model
        if err.sim_provider:
            dev.sim_provider = err.sim_provider
        if err.activation_plan:
            dev.activation_plan = err.activation_plan
        if err.firmware:
            dev.firmware = err.firmware
        if err.flash:
            dev.flash = err.flash
        if stype:
            dev.sim_type = stype
        merged_extra = {**(dev.extra or {}), **extra}
        dev.extra = merged_extra or None
        if dev.deleted_at:
            dev.deleted_at = dev.deleted_by = dev.delete_reason = None
    else:
        dev = Device(
            customer_id=cust.id,
            imei=final_imei,
            device_sr_no=err.device_sr_no,
            iccid=err.iccid,
            msisdn1=err.msisdn1,
            msisdn2=err.msisdn2,
            act_date=err.act_date,
            expiry_date=err.expiry_date,
            dispatched_date=err.dispatched_date,
            model=err.model,
            sim_provider=err.sim_provider,
            activation_plan=err.activation_plan,
            firmware=err.firmware,
            flash=err.flash,
            sim_type=stype,
            extra=extra or None,
            added_by=user.id if user else None
        )
        db.add(dev)
    db.delete(err)
    db.commit()
    db.refresh(dev)

    row_mirror = {
        "imei": dev.imei,
        "customer_name": cust.name,
        "device_sr_no": dev.device_sr_no,
        "iccid": dev.iccid,
        "msisdn1": dev.msisdn1,
        "msisdn2": dev.msisdn2,
        "act_date": dev.act_date,
        "expiry_date": dev.expiry_date,
        "dispatched_date": dev.dispatched_date,
        "model": dev.model,
        "sim_provider": dev.sim_provider,
        "activation_plan": dev.activation_plan,
        "firmware": dev.firmware,
        "flash": dev.flash,
        "sim_type": dev.sim_type,
        "extra": dev.extra or {}
    }
    try:
        if getattr(getattr(db, "bind", None), "dialect", None) and getattr(db.bind.dialect, "name", "") == "sqlite":
            customer_tables.mirror_rows([(cust.id, cust.name, row_mirror)], conn=db.connection())
            db.commit()
        else:
            customer_tables.mirror_rows([(cust.id, cust.name, row_mirror)])
    except Exception as me:
        print(f"Customer table mirror notice: {me}")

    log_activity(db, "error_force_pushed", f"Pushed errored entry #{error_id} (IMEI {final_imei}) to customer '{cust.name}'", user=user)
    db.commit()

    return {"success": True, "device_id": dev.id, "imei": final_imei, "customer_name": cust.name}


@router.post("/errors/bulk-push", dependencies=[Depends(require_permission("upload"))])
def bulk_push_errors_to_customer(data: schemas.BulkPushErrorsIn, db: Session = Depends(get_db),
                                 user: models.User = Depends(get_current_user)):
    """Bulk pushes errored entries directly into a specified customer."""
    query = db.query(models.ErroredRecord)
    if data.error_ids:
        query = query.filter(models.ErroredRecord.id.in_(data.error_ids))
    records = query.all()
    if not records:
        return {"pushed_count": 0, "customer_name": data.customer_name or "General_Customer"}

    target_cname = (data.customer_name or "").strip() or "General_Customer"
    cust = customer_by_name(db, target_cname)

    mirror_items = []
    pushed = 0
    for err in records:
        raw_imei = clean_imei(err.imei) if err.imei else ""
        extra = dict(err.extra or {})
        extra["pushed_with_error"] = True
        if err.error_reason:
            extra["original_error"] = err.error_reason

        if not raw_imei or len(raw_imei) < 5:
            alt = clean_imei(err.device_sr_no or err.iccid or "")
            if alt and len(alt) >= 5:
                raw_imei = alt[:20]
            else:
                raw_imei = f"INV_{uuid.uuid4().hex[:10].upper()}"

        final_imei = raw_imei[:20]
        stype = err.sim_type or ("E-SIM" if err.iccid else "Plastic SIM")
        existing_dev = db.query(Device).filter(Device.imei == final_imei).first()
        if existing_dev:
            dev = existing_dev
            dev.customer_id = cust.id
            if err.device_sr_no:
                dev.device_sr_no = err.device_sr_no
            if err.iccid:
                dev.iccid = err.iccid
            if err.msisdn1:
                dev.msisdn1 = err.msisdn1
            if err.msisdn2:
                dev.msisdn2 = err.msisdn2
            if err.act_date:
                dev.act_date = err.act_date
            if err.expiry_date:
                dev.expiry_date = err.expiry_date
            if err.dispatched_date:
                dev.dispatched_date = err.dispatched_date
            if err.model:
                dev.model = err.model
            if err.sim_provider:
                dev.sim_provider = err.sim_provider
            if err.activation_plan:
                dev.activation_plan = err.activation_plan
            if err.firmware:
                dev.firmware = err.firmware
            if err.flash:
                dev.flash = err.flash
            if stype:
                dev.sim_type = stype
            merged_extra = {**(dev.extra or {}), **extra}
            dev.extra = merged_extra or None
            if dev.deleted_at:
                dev.deleted_at = dev.deleted_by = dev.delete_reason = None
        else:
            dev = Device(
                customer_id=cust.id,
                imei=final_imei,
                device_sr_no=err.device_sr_no,
                iccid=err.iccid,
                msisdn1=err.msisdn1,
                msisdn2=err.msisdn2,
                act_date=err.act_date,
                expiry_date=err.expiry_date,
                dispatched_date=err.dispatched_date,
                model=err.model,
                sim_provider=err.sim_provider,
                activation_plan=err.activation_plan,
                firmware=err.firmware,
                flash=err.flash,
                sim_type=stype,
                extra=extra or None,
                added_by=user.id if user else None
            )
            db.add(dev)
        db.delete(err)
        pushed += 1
        mirror_items.append((cust.id, cust.name, {
            "imei": final_imei, "customer_name": cust.name, "device_sr_no": dev.device_sr_no,
            "iccid": dev.iccid, "msisdn1": dev.msisdn1, "msisdn2": dev.msisdn2,
            "act_date": dev.act_date, "expiry_date": dev.expiry_date, "dispatched_date": dev.dispatched_date,
            "model": dev.model, "sim_provider": dev.sim_provider, "activation_plan": dev.activation_plan,
            "firmware": dev.firmware, "flash": dev.flash, "sim_type": dev.sim_type, "extra": extra
        }))

    db.commit()
    if mirror_items:
        try:
            if getattr(getattr(db, "bind", None), "dialect", None) and getattr(db.bind.dialect, "name", "") == "sqlite":
                customer_tables.mirror_rows(mirror_items, conn=db.connection())
                db.commit()
            else:
                customer_tables.mirror_rows(mirror_items)
        except Exception as me:
            print(f"Customer table mirror notice: {me}")

    log_activity(db, "error_bulk_force_pushed", f"Bulk pushed {pushed} errored entries to customer '{cust.name}'", user=user)
    db.commit()

    return {"pushed_count": pushed, "customer_name": cust.name}


@router.delete("/errors/{error_id}", dependencies=[Depends(require_staff), Depends(require_permission("delete_invalid"))])
def move_error_to_cart(error_id: int, permanent: bool = False, db: Session = Depends(get_db),
                       user: models.User = Depends(get_current_user)):
    err = db.query(models.ErroredRecord).filter(models.ErroredRecord.id == error_id).first()
    if not err:
        raise HTTPException(404, "Errored record not found.")

    if permanent:
        db.delete(err)
        db.commit()
        log_activity(db, "error_permanent_delete", f"Permanently deleted errored entry #{error_id} from database", user=user)
        return {"deleted_forever": error_id, "permanent": True}

    # Find or create customer for this record so it can be held in Cart
    cname = (err.customer_name or "").strip() or "Unassigned Errors"
    cust = customer_by_name(db, cname)

    imei_val = (err.imei or "").strip()
    if not imei_val or len(imei_val) < 5 or db.query(Device).filter(Device.imei == imei_val).first():
        imei_val = f"ERR-{err.id}-{int(time.time())}"

    # Move to Cart (deleted_at set)
    dev = Device(
        customer_id=cust.id,
        imei=imei_val,
        device_sr_no=err.device_sr_no,
        iccid=err.iccid,
        msisdn1=err.msisdn1,
        msisdn2=err.msisdn2,
        act_date=err.act_date,
        expiry_date=err.expiry_date,
        dispatched_date=err.dispatched_date,
        model=err.model,
        sim_provider=err.sim_provider,
        activation_plan=err.activation_plan,
        firmware=err.firmware,
        flash=err.flash,
        sim_type=err.sim_type,
        extra=err.extra,
        added_by=user.id,
        deleted_at=datetime.utcnow(),
        deleted_by=user.user_id,
        delete_reason=f"Moved from Errored Entries: {err.error_reason[:150]}"
    )
    db.add(dev)
    db.delete(err)
    db.commit()

    log_activity(db, "error_trash", f"Errored entry #{error_id} moved to Cart", user=user)
    return {"moved_to_cart": error_id}


@router.post("/errors/empty", dependencies=[Depends(require_staff), Depends(require_permission("delete_invalid"))])
def move_all_errors_to_cart(limit: int = 0, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    total_remaining = db.query(models.ErroredRecord).count()
    query = db.query(models.ErroredRecord)
    if limit > 0:
        errors = query.limit(limit).all()
    else:
        errors = query.all()
    count = 0
    cust = customer_by_name(db, "Unassigned Errors")
    for err in errors:
        cname = (err.customer_name or "").strip()
        c = customer_by_name(db, cname) if cname else cust
        imei_val = (err.imei or "").strip()
        if not imei_val or len(imei_val) < 5 or db.query(Device).filter(Device.imei == imei_val).first():
            imei_val = f"ERR-{err.id}-{int(time.time())}"
        dev = Device(
            customer_id=c.id,
            imei=imei_val[:20],
            device_sr_no=str(err.device_sr_no or "")[:50] if err.device_sr_no else None,
            iccid=str(err.iccid or "")[:30] if err.iccid else None,
            msisdn1=str(err.msisdn1 or "")[:20] if err.msisdn1 else None,
            msisdn2=str(err.msisdn2 or "")[:20] if err.msisdn2 else None,
            act_date=err.act_date,
            expiry_date=err.expiry_date,
            dispatched_date=err.dispatched_date,
            model=str(err.model or "")[:100] if err.model else None,
            sim_provider=str(err.sim_provider or "")[:100] if err.sim_provider else None,
            activation_plan=str(err.activation_plan or "")[:100] if err.activation_plan else None,
            firmware=str(err.firmware or "")[:100] if err.firmware else None,
            flash=str(err.flash or "")[:100] if err.flash else None,
            sim_type=str(err.sim_type or "")[:20] if err.sim_type else None,
            extra=err.extra,
            added_by=user.id,
            deleted_at=datetime.utcnow(),
            deleted_by=user.user_id,
            delete_reason=f"Moved from Errored Entries: {err.error_reason[:150]}"
        )
        db.add(dev)
        db.delete(err)
        count += 1
    db.commit()
    remaining = max(0, total_remaining - count)
    log_activity(db, "errors_trash_all", f"Moved {count} errored entries to Cart", user=user)
    return {"moved_count": count, "remaining_count": remaining, "moved_all_to_cart": count}


@router.delete("/errors/permanent", dependencies=[Depends(require_staff), Depends(require_permission("delete_invalid"))])
def direct_delete_all_errors_permanent(limit: int = 0, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Directly and permanently deletes errored entries from the database forever."""
    total_remaining = db.query(models.ErroredRecord).count()
    if limit > 0:
        err_ids = [r.id for r in db.query(models.ErroredRecord.id).order_by(models.ErroredRecord.id.asc()).limit(limit).all()]
        if err_ids:
            count = db.query(models.ErroredRecord).filter(models.ErroredRecord.id.in_(err_ids)).delete(synchronize_session=False)
        else:
            count = 0
        db.commit()
        remaining = max(0, total_remaining - count)
        log_activity(db, "errors_permanent_purge", f"Permanently deleted {count} errored entries from database", user=user)
        return {"deleted_count": count, "remaining_count": remaining, "deleted_forever": count}
    else:
        count = db.query(models.ErroredRecord).delete(synchronize_session=False)
        db.commit()
        log_activity(db, "errors_permanent_purge", f"Permanently deleted all {count} errored entries from database", user=user)
        return {"deleted_count": count, "remaining_count": 0, "deleted_forever": count}


@router.delete("/errors/{error_id}/permanent", dependencies=[Depends(require_staff), Depends(require_permission("delete_invalid"))])
def direct_delete_single_error_permanent(error_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Directly and permanently deletes a single errored entry from the database forever."""
    err = db.query(models.ErroredRecord).filter(models.ErroredRecord.id == error_id).first()
    if not err:
        raise HTTPException(404, "Errored record not found.")
    db.delete(err)
    db.commit()
    log_activity(db, "error_permanent_delete", f"Permanently deleted errored entry #{error_id} from database", user=user)
    return {"deleted_forever": error_id}


@router.post("/errors/batch-delete", dependencies=[Depends(require_staff), Depends(require_permission("delete_invalid"))])
def direct_delete_batch_errors(data: schemas.BatchDeleteErrorsIn, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Directly and permanently deletes specific errored records by ID from the database forever."""
    if not data.error_ids:
        return {"deleted_count": 0, "deleted_forever": 0}
    count = db.query(models.ErroredRecord).filter(models.ErroredRecord.id.in_(data.error_ids)).delete(synchronize_session=False)
    db.commit()
    log_activity(db, "errors_batch_delete", f"Permanently deleted {count} errored entries from database", user=user)
    return {"deleted_count": count, "deleted_forever": count}


# ---------- Bulk Deletion via Excel matching ----------
@router.post("/devices/bulk-delete/preview", dependencies=[Depends(require_staff), Depends(require_permission("delete"))])
def bulk_delete_preview(data: schemas.BulkDeletePreviewIn, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Analyzes an uploaded Excel list of identifiers and previews matches in valid inventory."""
    col_name = (data.column or "imei").strip().lower()
    raw_vals = [str(v).strip() for v in data.values if v is not None and str(v).strip()]
    total_vals = len(raw_vals)
    if not total_vals:
        return {
            "total_values": 0, "unique_values": 0, "duplicate_values": 0,
            "matching_records": 0, "unmatched_values": 0, "eligible_for_deletion": 0,
            "column": col_name, "sample_matches": []
        }

    unique_vals = list(dict.fromkeys(raw_vals))
    dup_count = total_vals - len(unique_vals)

    col_map = {
        "imei": Device.imei,
        "iccid": Device.iccid,
        "device_sr_no": Device.device_sr_no,
        "msisdn1": Device.msisdn1,
        "customer_name": None
    }
    target_col = col_map.get(col_name, Device.imei)

    matched_ids = []
    matched_sample = []
    matched_val_set = set()
    chunk_size = 500

    for i in range(0, len(unique_vals), chunk_size):
        chunk = unique_vals[i:i + chunk_size]
        if target_col is not None:
            rows = db.query(Device.id, target_col).filter(
                Device.deleted_at.is_(None),
                target_col.in_(chunk)
            ).all()
            for dev_id, val in rows:
                matched_ids.append(dev_id)
                if val:
                    matched_val_set.add(str(val).strip().lower())
                if len(matched_sample) < 10:
                    matched_sample.append({"id": dev_id, "matched_value": val})
        else:
            rows = db.query(Device.id, Customer.name).join(Customer, Customer.id == Device.customer_id).filter(
                Device.deleted_at.is_(None),
                Customer.name.in_(chunk)
            ).all()
            for dev_id, val in rows:
                matched_ids.append(dev_id)
                if val:
                    matched_val_set.add(str(val).strip().lower())
                if len(matched_sample) < 10:
                    matched_sample.append({"id": dev_id, "matched_value": val})

    unique_lower = {v.lower() for v in unique_vals}
    unmatched_count = len(unique_lower - matched_val_set)

    return {
        "total_values": total_vals,
        "unique_values": len(unique_vals),
        "duplicate_values": dup_count,
        "matching_records": len(matched_ids),
        "unmatched_values": max(0, unmatched_count),
        "eligible_for_deletion": len(matched_ids),
        "column": col_name,
        "sample_matches": matched_sample
    }


@router.post("/devices/bulk-delete/execute", dependencies=[Depends(require_staff), Depends(require_permission("delete"))])
def bulk_delete_execute(data: schemas.BulkDeleteExecuteIn, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Executes bulk deletion of inventory based on uploaded Excel column matching."""
    col_name = (data.column or "imei").strip().lower()
    raw_vals = [str(v).strip() for v in data.values if v is not None and str(v).strip()]
    if not raw_vals:
        return {"success": True, "deleted_count": 0, "skipped_count": 0, "unmatched_count": 0, "failures_count": 0, "permanent": data.permanent, "column": col_name}

    unique_vals = list(dict.fromkeys(raw_vals))

    col_map = {
        "imei": Device.imei,
        "iccid": Device.iccid,
        "device_sr_no": Device.device_sr_no,
        "msisdn1": Device.msisdn1,
        "customer_name": None
    }
    target_col = col_map.get(col_name, Device.imei)

    deleted_count = 0
    matched_val_set = set()
    chunk_size = 500

    for i in range(0, len(unique_vals), chunk_size):
        chunk = unique_vals[i:i + chunk_size]
        if target_col is not None:
            dev_objs = db.query(Device).options(joinedload(Device.customer)).filter(
                Device.deleted_at.is_(None),
                target_col.in_(chunk)
            ).all()
        else:
            dev_objs = db.query(Device).options(joinedload(Device.customer)).join(Customer, Customer.id == Device.customer_id).filter(
                Device.deleted_at.is_(None),
                Customer.name.in_(chunk)
            ).all()

        if not dev_objs:
            continue

        for d in dev_objs:
            val = getattr(d, col_name, None) if col_name != "customer_name" else (d.customer.name if d.customer else None)
            if val:
                matched_val_set.add(str(val).strip().lower())

            cid, imei = d.customer_id, d.imei
            cname = d.customer.name if d.customer else ""

            if data.permanent:
                try:
                    customer_tables.delete_device_from_table(cid, imei, cname)
                except Exception:
                    pass
                db.delete(d)
            else:
                trash(d, user, data.reason or f"Bulk deleted via Excel matching {col_name}")
                try:
                    customer_tables.delete_device_from_table(cid, imei, cname)
                except Exception:
                    pass

            deleted_count += 1

        db.commit()

    unique_lower = {v.lower() for v in unique_vals}
    unmatched_count = max(0, len(unique_lower - matched_val_set))

    action_tag = "device_bulk_permanent_delete" if data.permanent else "device_bulk_trash"
    log_activity(
        db,
        action_tag,
        f"Bulk deleted {deleted_count} records via Excel (Column: {col_name}, Permanent: {data.permanent})",
        user=user
    )

    return {
        "success": True,
        "deleted_count": deleted_count,
        "skipped_count": 0,
        "unmatched_count": unmatched_count,
        "failures_count": 0,
        "permanent": data.permanent,
        "column": col_name
    }


# ---------- Bulk Edit by IMEI ----------
@router.post("/devices/bulk-edit/preview", dependencies=[Depends(require_staff), Depends(require_permission("edit"))])
def bulk_edit_preview(data: schemas.BulkEditPreviewIn, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Previews bulk editing of inventory records matched by IMEI."""
    input_rows = data.file_rows if data.file_rows is not None else (data.rows or [])
    total_rows = len(input_rows)
    imei_col_raw = (data.imei_col or data.imei_column or "imei").strip()
    mappings = data.field_mappings or {}

    if not total_rows:
        return {
            "total_rows": 0, "total_values": 0, "matching_imeis": 0, "matching_records": 0,
            "unmatched_imeis": 0, "unmatched_values": 0, "duplicates_in_file": 0, "duplicate_imeis": 0,
            "invalid_imeis": 0, "eligible_to_update": 0, "eligible_for_update": 0,
            "sample_changes": []
        }

    raw_imeis = []
    invalid_count = 0
    valid_row_map = {}
    seen_imeis = set()
    duplicate_count = 0

    for r in input_rows:
        # Extract row dictionary mapped to db fields
        if mappings:
            row_dict = {}
            for ex_col, db_fld in mappings.items():
                if ex_col in r:
                    row_dict[db_fld] = r[ex_col]
            raw_imei = str(r.get(imei_col_raw) or r.get("imei") or r.get(imei_col_raw.lower()) or "").strip()
        else:
            row_dict = dict(r)
            raw_imei = str(r.get(imei_col_raw) or r.get("imei") or r.get(imei_col_raw.lower()) or "").strip()

        if re.search(r"e\+", raw_imei, re.I):
            try:
                raw_imei = str(int(float(raw_imei)))
            except Exception:
                pass
        clean = re.sub(r"[^0-9A-Za-z_-]", "", raw_imei).strip()
        if not clean or len(clean) < 5:
            invalid_count += 1
            continue
        if clean in seen_imeis:
            duplicate_count += 1
        else:
            seen_imeis.add(clean)
            raw_imeis.append(clean)
            valid_row_map[clean] = row_dict

    matching_devices = {}
    chunk_size = 500
    for i in range(0, len(raw_imeis), chunk_size):
        chunk = raw_imeis[i:i + chunk_size]
        devs = db.query(Device).options(joinedload(Device.customer)).filter(
            Device.deleted_at.is_(None),
            Device.imei.in_(chunk)
        ).all()
        for d in devs:
            matching_devices[d.imei] = d

    matching_count = len(matching_devices)
    unmatched_count = max(0, len(raw_imeis) - matching_count)
    eligible_count = matching_count

    sample_changes_grouped = []
    sample_changes_flat = []
    for imei, d in list(matching_devices.items())[:25]:
        new_row = valid_row_map.get(imei, {})
        changes = {}
        for fld in ["model", "sim_provider", "activation_plan", "firmware", "flash", "sim_type", "device_sr_no", "iccid", "msisdn1", "msisdn2"]:
            new_val = new_row.get(fld)
            if new_val is not None and str(new_val).strip() != "":
                old_val = getattr(d, fld, None)
                if str(new_val).strip() != str(old_val or "").strip():
                    changes[fld] = {"old": str(old_val or ""), "new": str(new_val).strip()}

        for fld in ["act_date", "expiry_date", "dispatched_date"]:
            new_val = new_row.get(fld)
            if new_val is not None and str(new_val).strip() != "":
                old_val = getattr(d, fld, None)
                parsed_val = normalize_date_input(str(new_val).strip())
                if parsed_val and str(old_val or "") != str(parsed_val):
                    changes[fld] = {"old": str(old_val or ""), "new": str(parsed_val)}

        new_cust = new_row.get("customer_name")
        if new_cust and str(new_cust).strip():
            old_cust = d.customer.name if d.customer else ""
            if new_cust.strip().lower() != old_cust.lower():
                changes["customer_name"] = {"old": old_cust, "new": new_cust.strip()}

        for fld, diff in changes.items():
            sample_changes_flat.append({
                "imei": imei,
                "field": fld,
                "old_val": diff["old"],
                "new_val": diff["new"]
            })

        sample_changes_grouped.append({
            "id": d.id,
            "imei": imei,
            "customer_name": d.customer.name if d.customer else "",
            "changes": changes,
            "fields_changed_count": len(changes)
        })

    return {
        "total_rows": total_rows,
        "total_values": total_rows,
        "matching_imeis": matching_count,
        "matching_records": matching_count,
        "unmatched_imeis": unmatched_count,
        "unmatched_values": unmatched_count,
        "duplicates_in_file": duplicate_count,
        "duplicate_imeis": duplicate_count,
        "invalid_imeis": invalid_count,
        "eligible_to_update": eligible_count,
        "eligible_for_update": eligible_count,
        "sample_changes": sample_changes_flat if mappings else sample_changes_grouped,
        "detailed_samples": sample_changes_grouped
    }


@router.post("/devices/bulk-edit/execute", dependencies=[Depends(require_staff), Depends(require_permission("edit"))])
def bulk_edit_execute(data: schemas.BulkEditExecuteIn, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Executes bulk update of matching inventory records by IMEI."""
    input_rows = data.file_rows if data.file_rows is not None else (data.rows or [])
    total_input = len(input_rows)
    imei_col_raw = (data.imei_col or data.imei_column or "imei").strip()
    mappings = data.field_mappings or {}

    if not total_input:
        return {"success": True, "total_submitted": 0, "total_processed": 0, "updated": 0, "updated_count": 0, "skipped": 0, "skipped_count": 0, "unmatched": 0, "unmatched_count": 0, "duplicate_count": 0, "failures_count": 0, "errors": []}

    row_map = {}
    seen = set()
    dup_count = 0
    invalid_count = 0

    for r in input_rows:
        if mappings:
            row_dict = {}
            for ex_col, db_fld in mappings.items():
                if ex_col in r:
                    row_dict[db_fld] = r[ex_col]
            raw_imei = str(r.get(imei_col_raw) or r.get("imei") or r.get(imei_col_raw.lower()) or "").strip()
        else:
            row_dict = dict(r)
            raw_imei = str(r.get(imei_col_raw) or r.get("imei") or r.get(imei_col_raw.lower()) or "").strip()

        if re.search(r"e\+", raw_imei, re.I):
            try:
                raw_imei = str(int(float(raw_imei)))
            except Exception:
                pass
        clean = re.sub(r"[^0-9A-Za-z_-]", "", raw_imei).strip()
        if not clean or len(clean) < 5:
            invalid_count += 1
            continue
        if clean in seen:
            dup_count += 1
        else:
            seen.add(clean)
            row_map[clean] = row_dict

    unique_imeis = list(row_map.keys())
    updated_count = 0
    matched_set = set()
    failures_count = 0
    chunk_size = 500

    cust_cache = {c.name.lower(): c.id for c in db.query(Customer.id, Customer.name).filter(Customer.deleted_at.is_(None))}

    for i in range(0, len(unique_imeis), chunk_size):
        chunk = unique_imeis[i:i + chunk_size]
        devs = db.query(Device).options(joinedload(Device.customer)).filter(
            Device.deleted_at.is_(None),
            Device.imei.in_(chunk)
        ).all()

        for d in devs:
            matched_set.add(d.imei)
            new_data = row_map.get(d.imei, {})
            try:
                has_change = False
                for fld in ["model", "sim_provider", "activation_plan", "firmware", "flash", "sim_type", "device_sr_no", "iccid", "msisdn1", "msisdn2"]:
                    if fld in new_data:
                        v = new_data[fld]
                        if v is not None and str(v).strip() != "":
                            setattr(d, fld, str(v).strip())
                            has_change = True

                for fld in ["act_date", "expiry_date", "dispatched_date"]:
                    if fld in new_data:
                        v = new_data[fld]
                        if v is not None and str(v).strip() != "":
                            parsed_d = normalize_date_input(str(v).strip())
                            if parsed_d:
                                setattr(d, fld, parsed_d)
                                has_change = True

                if "customer_name" in new_data and str(new_data["customer_name"]).strip():
                    new_cname = str(new_data["customer_name"]).strip()
                    lower_c = new_cname.lower()
                    if lower_c not in cust_cache:
                        new_cust_obj = Customer(name=new_cname)
                        db.add(new_cust_obj)
                        db.flush()
                        cust_cache[lower_c] = new_cust_obj.id
                    d.customer_id = cust_cache[lower_c]
                    has_change = True

                if "extra" in new_data and isinstance(new_data["extra"], dict):
                    cur_extra = dict(d.extra or {})
                    for ek, ev in new_data["extra"].items():
                        if ev is not None and str(ev).strip() != "":
                            cur_extra[str(ek).strip()] = str(ev).strip()
                    d.extra = cur_extra
                    has_change = True

                if has_change:
                    updated_count += 1
            except Exception:
                failures_count += 1

        db.commit()

    unmatched_count = max(0, len(unique_imeis) - len(matched_set))
    log_activity(db, "bulk_edit", f"Bulk updated {updated_count} devices by IMEI (processed {total_input})", user=user)

    return {
        "success": True,
        "total_submitted": total_input,
        "total_processed": total_input,
        "updated": updated_count,
        "updated_count": updated_count,
        "unmatched": unmatched_count,
        "unmatched_count": unmatched_count,
        "skipped": dup_count + invalid_count,
        "skipped_count": dup_count + invalid_count,
        "duplicate_count": dup_count,
        "invalid_count": invalid_count,
        "failures_count": failures_count,
        "errors": []
    }

